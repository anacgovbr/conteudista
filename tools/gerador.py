# -*- coding: utf-8 -*-
# GERADOR PADRÃO DE CONTEÚDO v3.9 (imagens em base64; temas: govbr [padrão] e classico; narração das notas no HTML e no PPTX;
#   leitura acessível: Verdana / Atkinson Hyperlegible Next / OpenDyslexic, modo leitura dos slides, EPUB) — Apostila (HTML + DOCX + EPUB)
#   e Apresentação (HTML + PPTX) a partir de Markdown.
# Requisitos: Python 3 + python-docx + python-pptx + lxml (+ Pillow). HTML usa só biblioteca padrão.
# Fontes de leitura: tools/fontes/ (WOFF2 + OFL + manifesto.json). Sem as fontes, o conteúdo cai para Verdana e a geração continua.
# Uso: python gerador.py apostila.md apresentacao.md --prefixo NOME --rodape "TEXTO" [--barra "TEXTO"] [--saida pasta] [--sem-embutir] [--tema govbr|classico] [--audios pasta] [--avancar]
# Imagens ![alt](caminho|URL) sao convertidas para data:URI base64 (texto puro) nos .md de saida.
# Saída: Apostila_<prefixo>.md/.html/.docx/.epub/_OpenDyslexic.epub  e  Apresentacao_<prefixo>.md/.html/.pptx
import re,html,argparse,os,base64,io,mimetypes,zipfile,hashlib,uuid,datetime
from urllib.parse import unquote,unquote_to_bytes
E=html.escape
# ---------------- Avisos de acessibilidade (imagem sem descrição, fonte ausente etc.) ----------------
AVISOS=[]
def aviso(m):
    if m not in AVISOS:AVISOS.append(m)
# ---------------- Fontes de leitura (embutidas; sem CDN) ----------------
FONTES_DIR=os.path.join(os.path.dirname(os.path.abspath(__file__)),'fontes')
LEITURAS={'atkinson':('Atkinson Hyperlegible Next','atkinson-hyperlegible-next','OFL-Atkinson-Hyperlegible-Next.txt'),
          'opendyslexic':('OpenDyslexic','opendyslexic','OFL-OpenDyslexic.txt')}
FACES=((400,'normal',''),(400,'italic','-italic'),(700,'normal',''),(700,'italic','-italic'))
PILHA_LEITURA='Verdana,Tahoma,Arial,sans-serif'
def fontes_faces(fonte):
    "Faces existentes da familia: [(peso, estilo, nome_do_arquivo, caminho)]. Face ausente gera aviso e e ignorada."
    nome,base,_=LEITURAS[fonte];r=[];faltam=[]
    for peso,estilo,suf in FACES:
        arq='%s-%d%s.woff2'%(base,peso,suf);cam=os.path.join(FONTES_DIR,arq)
        if os.path.isfile(cam):r.append((peso,estilo,arq,cam))
        else:faltam.append(arq)
    if faltam:aviso('Fonte de leitura %s: %d de 4 arquivos ausentes em tools/fontes (o texto usa %s)'%(nome,len(faltam),'Verdana' if len(faltam)==4 else 'a face disponível'))
    return r
_FONTES_CSS=None
def fontes_css():
    "@font-face em base64 (Atkinson Hyperlegible Next e OpenDyslexic) + variavel --leitura. Calculado uma vez por execucao."
    global _FONTES_CSS
    if _FONTES_CSS is None:
        o=[]
        for fonte,(nome,base,_) in LEITURAS.items():
            for peso,estilo,arq,cam in fontes_faces(fonte):
                o.append('@font-face{font-family:"%s";src:url(data:font/woff2;base64,%s) format("woff2");font-weight:%d;font-style:%s;font-display:swap}'
                         %(nome,base64.b64encode(open(cam,'rb').read()).decode(),peso,estilo))
        o.append(':root{--leitura:%s}html[data-leitura="padrao"]{--leitura:%s}'%(PILHA_LEITURA,PILHA_LEITURA))
        o.append('html[data-leitura="atkinson"]{--leitura:"Atkinson Hyperlegible Next",%s}html[data-leitura="opendyslexic"]{--leitura:"OpenDyslexic",%s}'%(PILHA_LEITURA,PILHA_LEITURA))
        _FONTES_CSS='\n'.join(o)+'\n'
    return _FONTES_CSS
# Aplica a leitura salva antes da pintura (evita troca visível de fonte). Preferência global, separada do tema.
LEITURA_HEAD='(function(){var l=null;try{l=localStorage.getItem("conteudista.leitura")}catch(e){}if(["padrao","atkinson","opendyslexic"].indexOf(l)<0)l="padrao";document.documentElement.setAttribute("data-leitura",l)})();'
LEITURA_SEL=('<label for="leitura" class="lt">Leitura</label><select id="leitura"><option value="padrao">Padrão (Verdana)</option>'
             '<option value="atkinson">Atkinson Hyperlegible Next</option><option value="opendyslexic">OpenDyslexic</option></select>')
LEITURA_JS="""(function(){var s=document.getElementById('leitura'),h=document.documentElement,V=['padrao','atkinson','opendyslexic'];if(!s)return;
s.value=V.indexOf(h.getAttribute('data-leitura'))>-1?h.getAttribute('data-leitura'):'padrao';
s.addEventListener('change',function(){var v=V.indexOf(s.value)>-1?s.value:'padrao';h.setAttribute('data-leitura',v);try{localStorage.setItem('conteudista.leitura',v)}catch(e){}
document.dispatchEvent(new CustomEvent('conteudista:leitura',{detail:v}))})})();"""
# ============================== PARSER COMUM ==============================
def slug(s):
    s=re.sub(r'<[^>]+>','',s).lower().translate(str.maketrans('áàâãäéèêëíìîïóòôõöúùûüçñ','aaaaaeeeeiiiiooooouuuucn'))
    return re.sub(r'[^a-z0-9]+','-',s).strip('-')
LI=re.compile(r'^(\s*)([-*]|\d+\.)\s+(.*)$')
CODIGO_RX=re.compile(r'(?ms)(^```[^\n]*\n.*?^```[ \t]*$|`[^`\n]+`)')
def sem_comentarios(md):
    """Remove comentarios HTML (<!-- -->) do texto, preservando os que aparecem em blocos de codigo e em codigo em linha."""
    out=[];p=0
    for m in CODIGO_RX.finditer(md):
        out.append(re.sub(r'<!--.*?-->','',md[p:m.start()],flags=re.S));out.append(m.group(0));p=m.end()
    out.append(re.sub(r'<!--.*?-->','',md[p:],flags=re.S))
    return ''.join(out)
def blocks(md):
    L=md.split('\n');n=len(L);i=0;o=[]
    st=lambda l:l.startswith(('#','```','|','>')) or LI.match(l) or l.strip() in('','---')
    while i<n:
        l=L[i]
        if not l.strip() or l.strip()=='---':i+=1;continue
        if l.startswith('```'):
            lg=l[3:].strip();b=[];i+=1
            while i<n and not L[i].startswith('```'):b.append(L[i]);i+=1
            i+=1;o.append(('code',lg,'\n'.join(b)));continue
        m=re.match(r'^(#{1,6})\s+(.*)$',l)
        if m:o.append(('h',len(m.group(1)),m.group(2).strip()));i+=1;continue
        m=re.match(r'^\s*!\[([^\]]*)\]\(([^)\s]+)\)\s*$',l)
        if m:o.append(('img',m.group(1),m.group(2)));i+=1;continue
        if l.startswith('|'):
            r=[]
            while i<n and L[i].startswith('|'):r.append(L[i]);i+=1
            cs=lambda x:[y.strip().replace('\\|','|') for y in re.split(r'(?<!\\)\|',x.strip().strip('|'))]
            rows=[cs(r[0])]+[cs(x) for x in r[2:]];o.append(('table',rows));continue
        if l.startswith('>'):
            b=[]
            while i<n and L[i].startswith('>'):b.append(L[i][1:].lstrip());i+=1
            k=''
            if b and b[0].startswith('[!') and ']' in b[0]:k=b[0][2:b[0].index(']')].lower();b[0]=b[0][b[0].index(']')+1:].strip()
            o.append(('quote',k,blocks('\n'.join(b))));continue
        if LI.match(l):
            it=[];base=None
            while i<n and (LI.match(L[i]) or (L[i].startswith('  ') and L[i].strip())):
                m=LI.match(L[i])
                if m:
                    ind=len(m.group(1));base=ind if base is None else base
                    t=m.group(3);ck=None;mm=re.match(r'^\[( |x)\]\s(.*)$',t)
                    if mm:ck=mm.group(1)=='x';t=mm.group(2)
                    it.append([min(3,max(0,(ind-base)//2)),m.group(2)[0].isdigit(),ck,t])
                elif it:it[-1][3]+=' '+L[i].strip()
                i+=1
            o.append(('list',it));continue
        b=[]
        while i<n and not st(L[i]):b.append(L[i].strip());i+=1
        if not b:b=[l.strip()];i+=1
        o.append(('p',' '.join(b)))
    return o
def runs(t):
    """Inline -> lista de (texto, negrito, itálico, tipo[0 texto|1 código|2 imagem], url)."""
    out=[];pat=re.compile(r'`([^`]+)`|!\[([^\]]*)\]\(([^)\s]+)\)|\*\*(.+?)\*\*|\[([^\]]+)\]\(([^)\s]+)\)|(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])')
    p=0
    for m in pat.finditer(t):
        if m.start()>p:out.append((t[p:m.start()],0,0,0,None))
        if m.group(1) is not None:out.append((m.group(1),0,0,1,None))
        elif m.group(3) is not None:out.append((m.group(2),0,0,2,m.group(3)))
        elif m.group(4) is not None:
            for x in runs(m.group(4)):out.append((x[0],1,x[2],x[3],x[4]))
        elif m.group(6) is not None:out.append((m.group(5),0,0,0,m.group(6)))
        else:out.append((m.group(7),0,1,0,None))
        p=m.end()
    if p<len(t):out.append((t[p:],0,0,0,None))
    return out
def plain(t):return ''.join(r[0] for r in runs(t) if r[3]!=2)
IMG=re.compile(r'!\[([^\]]*)\]\(([^)\s]+)\)')
def img_bytes(src):
    """Bytes de uma imagem data:URI ou arquivo local; None se indisponível ou SVG."""
    try:
        if src.startswith('data:'):
            h,d=src.split(',',1);mime=h[5:].split(';')[0]
            raw=base64.b64decode(d) if ';base64' in h else unquote_to_bytes(d)
        elif os.path.isfile(src):
            raw=open(src,'rb').read();mime=mimetypes.guess_type(src)[0] or ''
        else:return None
        if 'svg' in mime:return None
        return raw
    except Exception:return None
def img_size(raw):
    try:
        from PIL import Image
        return Image.open(io.BytesIO(raw)).size
    except Exception:return (800,450)
def embed_images(md,base_dir='.'):
    """Substitui imagens locais/remotas por data:URI base64. Retorna (md, embutidas, falhas)."""
    ok=[0];fail=[]
    def rep(m):
        alt,src=m.group(1),m.group(2)
        if src.startswith('data:'):return m.group(0)
        try:
            if re.match(r'https?://',src):
                import urllib.request
                with urllib.request.urlopen(src,timeout=15) as r:raw=r.read();mime=r.headers.get_content_type()
            else:
                p=unquote(src if os.path.isabs(src) else os.path.join(base_dir,src))
                raw=open(p,'rb').read();mime=mimetypes.guess_type(p)[0]
        except Exception:fail.append(src);return m.group(0)
        ok[0]+=1
        return '![%s](data:%s;base64,%s)'%(alt,mime or 'image/png',base64.b64encode(raw).decode())
    partes=[];p=0
    for m in CODIGO_RX.finditer(md):
        partes.append(IMG.sub(rep,md[p:m.start()]));partes.append(m.group(0));p=m.end()
    partes.append(IMG.sub(rep,md[p:]))
    return ''.join(partes),ok[0],fail
def mermaid_edges(code):
    lab={};rows=[]
    node=r'([A-Za-z0-9_]+)\s*(\(\(\(.*?\)\)\)|\(\(.*?\)\)|\[.*?\]|\{.*?\}|\(.*?\))?'
    def nm(i,s):
        if s:
            s=re.sub(r'^[\[\(\{]+|[\]\)\}]+$','',s);s=re.sub(r'<br\s*/?>',' — ',s).replace('**','');lab[i]=s.strip()
        return i
    for ln in code.split('\n'):
        ln=ln.strip()
        if '-->' not in ln:
            m=re.match(node+r'\s*$',ln)
            if m and m.group(2):nm(m.group(1),m.group(2))
            continue
        parts=re.split(r'\s*-->\s*',ln);prev=None
        for pt in parts:
            cond='';m=re.match(r'^\|([^|]*)\|\s*(.*)$',pt)
            if m:cond=m.group(1);pt=m.group(2)
            m=re.match(node,pt)
            if not m:continue
            cur=nm(m.group(1),m.group(2))
            if prev:rows.append([prev,cond,cur])
            prev=cur
    f=lambda x:lab.get(x,x) if lab.get(x,x) not in('O','**O**') else ('Início' if x.endswith('00') else 'Fim' if x.endswith('99') else x)
    return [['De','Condição','Para']]+[[f(a),c or '—',f(b)] for a,c,b in rows]
# ============================== HTML ==============================
def inl(t):
    s=''
    for x,b,i,c,u in runs(t):
        if c==2:s+='<img class="il" src="%s" alt="%s">'%(E(u),E(x));continue
        h=E(x,quote=False)
        if c:h='<code>'+h+'</code>'
        if i:h='<em>'+h+'</em>'
        if b:h='<strong>'+h+'</strong>'
        if u:h='<a href="%s" target="_blank" rel="noopener">%s</a>'%(E(u),h)
        s+=h
    return s
def lst_html(it):
    def rec(k,d):
        tag='ol' if it[k][1] else 'ul';o=''
        while k<len(it) and it[k][0]>=d:
            if it[k][0]>d:sub,k=rec(k,it[k][0]);o=o[:-5]+sub+'</li>';continue
            dp,od,ck,t=it[k]
            o+=('<li class="task"><label><input type="checkbox"%s> %s</label></li>'%(' checked' if ck else '',inl(t))) if ck is not None else '<li>'+inl(t)+'</li>';k+=1
        return '<%s>%s</%s>'%(tag,o,tag),k
    return rec(0,it[0][0])[0]
def html_of(bl,toc=None):
    o=[]
    for b in bl:
        k=b[0]
        if k=='h':
            sid=slug(b[2])
            if toc is not None and b[1] in(2,3):toc.append((b[1],sid,plain(b[2])))
            o.append('<h%d id="%s">%s</h%d>'%(b[1],sid,inl(b[2]),b[1]))
        elif k=='p':o.append('<p>'+inl(b[1])+'</p>')
        elif k=='img':
            if not b[1].strip():aviso('Imagem sem descrição (texto alternativo vazio): %s'%(b[2][:60]+'…' if len(b[2])>60 else b[2]))
            o.append('<figure><img src="%s" alt="%s">%s</figure>'%(E(b[2]),E(b[1]),'<figcaption>%s</figcaption>'%E(b[1]) if b[1] else ''))
        elif k=='table':
            r=b[1];t='<div class="tbl" role="region" aria-label="Tabela" tabindex="0"><table><thead><tr>'+''.join('<th scope="col">%s</th>'%inl(c) for c in r[0])+'</tr></thead><tbody>'
            t+=''.join('<tr>'+''.join('<td>%s</td>'%inl(c) for c in x)+'</tr>' for x in r[1:]);o.append(t+'</tbody></table></div>')
        elif k=='list':o.append(lst_html(b[1]))
        elif k=='quote':o.append('<blockquote class="%s">%s</blockquote>'%(b[1],html_of(b[2])))
        elif k=='code':o.append(_mermaid_html(b[2]) if b[1]=='mermaid' else '<pre class="code" tabindex="0"><code>%s</code></pre>'%E(b[2]))
    return '\n'.join(o)
def _tabela_fluxo(code,foco=True):
    "Alternativa textual do fluxo Mermaid: tabela De / Condição / Para."
    try:rows=mermaid_edges(code)
    except Exception:rows=None
    if not rows or len(rows)<2:return '<pre class="code"%s><code>%s</code></pre>'%(' tabindex="0"' if foco else '',E(code))
    t='<table><thead><tr>'+''.join('<th scope="col">%s</th>'%E(c) for c in rows[0])+'</tr></thead><tbody>'
    t+=''.join('<tr>'+''.join('<td>%s</td>'%E(c) for c in x)+'</tr>' for x in rows[1:])+'</tbody></table>'
    return '<div class="tbl" role="region" aria-label="Tabela do fluxo" tabindex="0">%s</div>'%t if foco else '<div class="tbl">%s</div>'%t
def _mermaid_html(code):
    # Desenho (Mermaid) + tabela do fluxo em <details> na tela + cópia só para impressão (não depende de JavaScript).
    return ('<div class="mw" role="img" aria-label="%s"><pre class="mermaid">%s</pre></div>'
            '<details class="fluxo-tabela"><summary>Tabela do fluxo</summary>%s</details>'
            '<div class="fluxo-print" aria-hidden="true">%s</div>')%(E(_fluxo_alt(code)),E(code),_tabela_fluxo(code),_tabela_fluxo(code,False))
MER='<script type="module">try{const m=(await import("https://cdn.jsdelivr.net/npm/mermaid@10/dist/mermaid.esm.min.mjs")).default;m.initialize({startOnLoad:false});await m.run({querySelector:"pre.mermaid"})}catch(e){console.warn("Mermaid offline",e)}</script>'
# ============================== TEMAS (layout dos arquivos gerados) ==============================
# Cada tema troca a paleta e a fonte de TODOS os arquivos (HTML, DOCX e PPTX). As cores "base" do código são as do
# tema clássico; o mapa de cada tema diz por qual cor cada uma é trocada. Para criar um tema novo (ex.: ANAC),
# copie o bloco 'govbr', ajuste 'map', fontes e 'extra_*' e registre em TEMAS.
_GOV_MAP={'#0b3d6e':'#0c326f','#0a7cc1':'#1351b4','#19a974':'#168821','#0a5c9a':'#155bcb','#eaf1f8':'#dbe8fb','#dfe5ee':'#cccccc','#d9e2ee':'#cccccc',
 '#1b2430':'#333333','#5a6676':'#636363','#f6f8fb':'#f8f8f8','#eef4fa':'#edf5ff','#0d1b2a':'#071d41','#1e7b4f':'#168821','#e9f6f0':'#e3f5e1',
 '#fff4e0':'#fff5c2','#9a5b00':'#8a5c00','#f1f4f8':'#f8f8f8','#f2f7fc':'#edf5ff','#d6ecff':'#dbe8fb','#e8f2fb':'#edf5ff','#cfe3f5':'#c5d4eb'}
TEMAS={
 'govbr':{'nome':'Padrão gov.br (Design System do Governo Federal)','map':_GOV_MAP,
   'fonte':"'Rawline','Raleway',\"Segoe UI\",Arial,sans-serif",
   'import':"@import url('https://fonts.cdnfonts.com/css/rawline');@import url('https://fonts.googleapis.com/css2?family=Raleway:wght@400;600;700;800&display=swap');",
   'extra_g':"""[data-theme=light] header{background:#fff;color:#333;border-top:4px solid #1351b4;box-shadow:0 1px 6px rgba(0,0,0,.16)}
[data-theme=light] header .t{color:#333}[data-theme=light] header #pg{color:#636363}
[data-theme=light] header button,[data-theme=light] header select{background:#fff;color:#1351b4;border:1px solid #1351b4;border-radius:100em;font-weight:600;padding:6px 14px}
[data-theme=light] header button:hover{background:rgba(19,81,180,.12)}
[data-theme=light] article{border:0;border-radius:4px;box-shadow:0 1px 6px rgba(0,0,0,.16)}
[data-theme=light] th{background:#ededed;color:#333}
h1,h2,h3{font-weight:700}.ab{border-radius:100em;font-weight:600;padding:8px 20px}
:focus-visible{outline:4px dashed #c2850c;outline-offset:4px}
nav a.on{font-weight:600}""",
   'extra_s':""".s h1{font-weight:700}.s.title h1{border-bottom-color:#ffcd07}.s.title:before{background:#ffcd07}
#hud button{border-radius:100em}""",
   'docx_font':'Arial','pptx_font':'Arial','pptx_destaque':'FFCD07'},
 'classico':{'nome':'Clássico (layout original do Conteudista)','map':{},'fonte':None,'import':'','extra_g':'','extra_s':'',
   'docx_font':'Calibri','pptx_font':'Segoe UI','pptx_destaque':None},
 'anac20-claro':{'nome':'ANAC 20 anos — claro','map':_GOV_MAP,'fonte':"'Rawline','Raleway',\"Segoe UI\",Arial,sans-serif",'import':'','extra_g':'','extra_s':'',
   'docx_font':'Arial','pptx_font':'Arial','pptx_destaque':'D5A900','pptx_modelo':'anac20-claro','pptx_escuro':False},
 'anac20-escuro':{'nome':'ANAC 20 anos — escuro','map':_GOV_MAP,'fonte':"'Rawline','Raleway',\"Segoe UI\",Arial,sans-serif",'import':'','extra_g':'','extra_s':'',
   'docx_font':'Arial','pptx_font':'Arial','pptx_destaque':'D5A900','pptx_modelo':'anac20-escuro','pptx_escuro':True},
}
TEMA_PADRAO='govbr'
def tema_de(t):return TEMAS.get(t or TEMA_PADRAO) or TEMAS[TEMA_PADRAO]
def _cor(h,T):
    h=h.lower();h=h if h.startswith('#') else '#'+h;return T['map'].get(h,h)
def tcss(css,T,extra=''):
    css=re.sub(r'#[0-9a-fA-F]{6}\b',lambda m:_cor(m.group(0),T),css)
    if T.get('fonte'):css=css.replace('"Segoe UI",Roboto,Arial,sans-serif',T['fonte'])
    return (T.get('import') or '')+css+(extra or '')
def thex(h,T):return _cor(h,T)[1:].upper()
def trgb(h,T):
    from docx.shared import RGBColor as _R
    return _R.from_string(thex(h,T))
CSS_G=""":root{--qi:#0a7cc1;--qib:#eaf1f8;--qt:#1e7b4f;--qtb:#e9f6f0;--qw:#9a5b00;--qwb:#fff4e0;--bg:#f6f8fb;--pa:#fff;--ink:#1b2430;--pri:#0b3d6e;--acc:#0a7cc1;--line:#dfe5ee;--th:#eaf1f8;--ok:#1e7b4f;--wa:#9a5b00;--code:#f1f4f8}
[data-theme=dark]{--qi:#5fb3ff;--qib:#16304d;--qt:#5fd39b;--qtb:#123326;--qw:#ffc266;--qwb:#3a2c10;--bg:#0f1520;--pa:#162030;--ink:#e6edf5;--pri:#8cc4ff;--acc:#5fb3ff;--line:#2a3a50;--th:#1e2c40;--ok:#5fd39b;--wa:#ffc266;--code:#1c283a}
[data-theme=contrast]{--qi:#00e5ff;--qib:#001f2b;--qt:#00ff88;--qtb:#002614;--qw:#ffd400;--qwb:#2b2400;--bg:#000;--pa:#000;--ink:#fff;--pri:#ffd400;--acc:#00e5ff;--line:#fff;--th:#1a1a1a;--ok:#0f8;--wa:#ffd400;--code:#111}
*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;background:var(--bg);color:var(--ink);font:16px/1.6 "Segoe UI",Roboto,Arial,sans-serif}
header{position:sticky;top:0;z-index:9;background:#0b3d6e;color:#fff;display:flex;align-items:center;gap:12px;padding:10px 18px}header .t{font-weight:700;flex:1;font-size:15px}
header button,header select{background:rgba(255,255,255,.15);color:#fff;border:1px solid rgba(255,255,255,.4);border-radius:6px;padding:6px 10px;font:inherit;font-size:13px;cursor:pointer}header option{color:#000}
.wrap{display:grid;grid-template-columns:290px 1fr;max-width:1400px;margin:auto}nav{position:sticky;top:52px;height:calc(100vh - 52px);overflow:auto;padding:18px 12px;border-right:1px solid var(--line);font-size:14px}
nav ul{list-style:none;margin:0;padding:0}nav .lv3{padding-left:14px;font-size:13px}nav a{display:block;padding:4px 8px;border-radius:6px;color:var(--ink);text-decoration:none}nav a:hover,nav a.on{background:var(--th);color:var(--pri)}
main{padding:28px 40px 80px;min-width:0}article{background:var(--pa);border:1px solid var(--line);border-radius:12px;padding:36px 48px;max-width:1000px}
h1{color:var(--pri);font-size:30px;line-height:1.25}h2{color:var(--pri);font-size:23px;margin:42px 0 12px;padding-top:6px;border-top:3px solid var(--acc);scroll-margin-top:70px}h3{color:var(--pri);font-size:18px;margin:28px 0 8px;scroll-margin-top:70px}a{color:var(--acc);word-break:break-word}
.tbl{overflow-x:auto;margin:14px 0}table{border-collapse:collapse;width:100%;font-size:14.5px}th,td{border:1px solid var(--line);padding:8px 10px;text-align:left;vertical-align:top}th{background:var(--th);color:var(--pri)}
code{background:var(--code);padding:1px 5px;border-radius:4px;font:.9em Consolas,monospace}pre.code{background:var(--code);padding:14px;border-radius:8px;overflow:auto}.mw{background:#fff;border:1px solid var(--line);border-radius:10px;padding:12px;margin:14px 0;overflow-x:auto}pre.mermaid{margin:0;color:#333}
figure{margin:16px 0;text-align:center}figure img{max-width:100%;height:auto;border:1px solid var(--line);border-radius:8px}figcaption{color:#5a6676;font-size:13px;margin-top:4px}img.il{max-height:1.6em;vertical-align:middle}
blockquote{margin:16px 0;padding:12px 16px;border-left:5px solid var(--acc);background:var(--th);border-radius:0 8px 8px 0}blockquote p{margin:4px 0}blockquote.info{border-left-color:var(--qi);background:var(--qib)}blockquote.tip{border-left-color:var(--qt);background:var(--qtb)}blockquote.warning{border-left-color:var(--qw);background:var(--qwb)}
li.task{list-style:none;margin-left:-22px}li.task input{transform:scale(1.15);margin-right:6px;accent-color:var(--acc)}.hid h3~*{filter:blur(6px);user-select:none;pointer-events:none}
.ab{margin:8px 0;background:var(--acc);color:#fff;border:0;border-radius:6px;padding:8px 14px;font:inherit;cursor:pointer}
@media(max-width:900px){.wrap{grid-template-columns:1fr}nav{display:none}main{padding:12px}article{padding:20px}}
@media print{header,nav,.ab{display:none}.wrap{display:block}article{border:0;padding:0;max-width:none}.hid h3~*{filter:none}h2{break-before:page}table,.mw,figure{break-inside:avoid}}"""
JS_G="""(function(){var K='__K__',th=document.getElementById('th'),v=localStorage.getItem(K+'t')||'light';document.documentElement.dataset.theme=v;th.value=v;
th.onchange=function(){document.documentElement.dataset.theme=th.value;localStorage.setItem(K+'t',th.value)};
var b=[].slice.call(document.querySelectorAll('li.task input')),s=JSON.parse(localStorage.getItem(K+'c')||'{}');
function p(){document.getElementById('pg').textContent='Checklists: '+b.filter(function(x){return x.checked}).length+'/'+b.length}
b.forEach(function(x,i){x.checked=!!s[i];x.onchange=function(){s[i]=x.checked;localStorage.setItem(K+'c',JSON.stringify(s));p()}});p();
document.getElementById('rs').onclick=function(){s={};localStorage.removeItem(K+'c');b.forEach(function(x){x.checked=false});p()};
var g=document.getElementById('gab');if(g){var a=document.createElement('button');a.type='button';a.className='ab';a.setAttribute('aria-controls','gab');g.before(a);
function sg(on){g.classList.toggle('hid',!on);if(on){g.removeAttribute('aria-hidden');g.removeAttribute('inert')}else{g.setAttribute('aria-hidden','true');g.setAttribute('inert','')}
a.setAttribute('aria-expanded',on?'true':'false');a.textContent=on?'Ocultar gabarito':'Mostrar gabarito'}sg(false);
a.onclick=function(){sg(g.classList.contains('hid'))}}
var l=[].slice.call(document.querySelectorAll('nav a')),h=l.map(function(x){return document.getElementById(x.getAttribute('href').slice(1))});
addEventListener('scroll',function(){var y=scrollY+90,k=0;h.forEach(function(e,i){if(e&&e.offsetTop<=y)k=i});l.forEach(function(x,i){x.classList.toggle('on',i===k)})})})();"""
CSS_S=""":root{--pri:#0b3d6e;--acc:#0a7cc1;--ac2:#19a974;--ink:#1b2430;--mu:#5a6676;--line:#d9e2ee;--soft:#eef4fa}*{box-sizing:border-box}
html,body{margin:0;height:100%;background:#0d1b2a;font-family:"Segoe UI",Roboto,Arial,sans-serif;overflow:hidden}#st{position:absolute;left:50%;top:50%;width:1280px;height:720px}
.s{position:absolute;inset:0;background:#fff;color:var(--ink);display:none;overflow:hidden}.s.on{display:block}
.s:before{content:"";position:absolute;left:0;top:0;bottom:0;width:14px;background:linear-gradient(var(--pri),var(--acc))}
.s:after{content:"__F__";position:absolute;left:60px;bottom:22px;font-size:13px;color:var(--mu)}.in{position:absolute;left:60px;right:60px;top:44px;bottom:60px;display:flex;flex-direction:column}
.s h1{font-size:40px;color:var(--pri);margin:0 0 22px;padding-bottom:12px;border-bottom:3px solid var(--acc);line-height:1.15}.s h2{font-size:26px;color:var(--acc);margin:0 0 14px}
.s p,.s li{font-size:23px;line-height:1.45;margin:6px 0}.s ul,.s ol{padding-left:34px;margin:4px 0}.s li li{font-size:20px}.s li::marker{color:var(--acc);font-weight:700}
.s table{border-collapse:collapse;width:100%;font-size:19px;margin:4px 0 10px}.s th{background:var(--pri);color:#fff;text-align:left;padding:9px 12px}.s td{border-bottom:1px solid var(--line);padding:8px 12px;vertical-align:top}.s tr:nth-child(even) td{background:var(--soft)}
.s a{color:var(--acc)}.s code{background:var(--soft);padding:1px 6px;border-radius:4px;color:var(--pri)}.s strong{color:var(--pri)}.s blockquote{margin:14px 0 0;padding:12px 18px;background:#e9f6f0;border-left:6px solid var(--ac2);border-radius:0 10px 10px 0}.s blockquote p{font-size:21px;margin:0}.s blockquote.info{background:#eaf1f8;border-left-color:#0a7cc1}.s blockquote.tip{background:#e9f6f0;border-left-color:#1e7b4f}.s blockquote.warning{background:#fff4e0;border-left-color:#9a5b00}
.s figure{margin:6px 0;text-align:center}.s figure img{max-width:100%;max-height:440px;object-fit:contain}.s figcaption{font-size:15px;color:var(--mu)}.s img.il{max-height:1.4em;vertical-align:middle}
.s .mw{border:1px solid var(--line);border-radius:10px;padding:8px;text-align:center}.s pre.mermaid{margin:0}.s .mw.fx{padding:6px}.s .mw img.fxi{display:block;margin:0 auto;max-width:100%;max-height:520px;width:auto;height:auto}body.mm .s:not(.on){display:block;visibility:hidden}
.s.title{background:linear-gradient(135deg,#0b3d6e,#0a5c9a 60%,#0a7cc1);color:#fff}.s.title:before{background:var(--ac2)}.s.title:after{color:#cfe3f5}.s.title .in{justify-content:center}
.s.title h1{color:#fff;font-size:58px;border-bottom:4px solid var(--ac2);display:inline-block}.s.title h2{color:#d6ecff;font-size:32px}.s.title p{color:#e8f2fb;font-size:22px}.s.title strong{color:#fff}
.s.highlight{background:linear-gradient(#fff,#f2f7fc)}.nt{display:none}.sn .s.on .nt{display:block;position:absolute;left:60px;right:60px;bottom:44px;background:#fff8d6;border:1px solid #e6cf6a;border-radius:8px;padding:10px 14px;font-size:16px}
#bar{position:fixed;left:0;bottom:0;height:5px;background:var(--ac2)}#hud{position:fixed;right:14px;bottom:12px;display:flex;gap:8px;align-items:center;font-size:13px;color:#cfe3f5;background:rgba(11,61,110,.9);padding:5px 8px;border-radius:10px;opacity:.35}#hud:hover,#hud:focus-within{opacity:1}
#hud button{background:rgba(255,255,255,.12);color:#fff;border:1px solid rgba(255,255,255,.3);border-radius:6px;padding:6px 10px;cursor:pointer}"""
JS_S="""(function(){var D=document,H=D.documentElement,B=D.body,S=[].slice.call(D.querySelectorAll('.s')),N=S.length,i=0,st=D.getElementById('st'),K='conteudista.apresentacao.modo';
function $(x){return D.getElementById(x)}var bLer=$('bLer'),b1=$('b1'),b2=$('b2');
function ler(){return B.classList.contains('modo-leitura')}
function f(){st.style.transform=ler()?'':'translate(-50%,-50%) scale('+Math.min(innerWidth/1280,innerHeight/720)*.98+')'}
function ajustar(s){if(!s)return;var c=s.querySelector('.conteudo'),n=s.querySelector('.in');if(!c||!n)return;
c.style.transform='';c.style.width='';c.style.transformOrigin='';n.classList.remove('rola');n.removeAttribute('tabindex');if(ler())return;
var a=n.clientHeight,h=c.scrollHeight;if(!a||h<=a+2)return;var k=a/h;if(k<.85){k=.85;n.classList.add('rola');n.setAttribute('tabindex','0')}
c.style.transformOrigin='top left';c.style.width=(100/k)+'%';c.style.transform='scale('+k+')'}
function marca(){S.forEach(function(x,k){var v=ler()||k===i;x.setAttribute('aria-hidden',v?'false':'true');if(v)x.removeAttribute('inert');else x.setAttribute('inert','')})}
function g(n){i=Math.max(0,Math.min(N-1,n));S.forEach(function(x,k){x.classList.toggle('on',k===i)});marca();
var c=$('c');if(c)c.textContent='Slide '+(i+1)+' de '+N;var b=$('bar');if(b)b.style.width=(i+1)/N*100+'%';
try{history.replaceState(null,'','#'+(i+1))}catch(e){}ajustar(S[i]);D.dispatchEvent(new CustomEvent('conteudista:slide',{detail:i}))}
function fs(){if(D.fullscreenElement)D.exitFullscreen();else if(H.requestFullscreen)H.requestFullscreen().catch(function(){})}
D.addEventListener('fullscreenchange',function(){if(b2)b2.setAttribute('aria-pressed',D.fullscreenElement?'true':'false')});
function nt(){var on=B.classList.toggle('sn');if(b1)b1.setAttribute('aria-pressed',on?'true':'false')}
function visivel(){for(var k=0;k<N;k++){if(S[k].getBoundingClientRect().bottom>90)return k}return i}
function modo(on,salvar){var era=ler();if(era&&!on)i=visivel();B.classList.toggle('modo-leitura',on);H.classList.toggle('ml',on);
if(bLer){bLer.setAttribute('aria-pressed',on?'true':'false');bLer.textContent=on?'Modo slides':'Modo leitura'}
[].slice.call(D.querySelectorAll('details.fluxo-tabela')).forEach(function(d){if(on){if(!d.open){d.open=true;d.setAttribute('data-auto','1')}}else if(d.getAttribute('data-auto')){d.open=false;d.removeAttribute('data-auto')}});
if(salvar){try{localStorage.setItem(K,on?'leitura':'slides')}catch(e){}}
f();g(i);if(on&&!era&&i>0)S[i].scrollIntoView();D.dispatchEvent(new CustomEvent('conteudista:modo',{detail:on?'leitura':'slides'}))}
function campo(t){var n=(t&&t.tagName||'').toLowerCase();return n=='input'||n=='select'||n=='textarea'||!!(t&&t.isContentEditable)}
addEventListener('keydown',function(e){if(e.ctrlKey||e.metaKey||e.altKey||e.defaultPrevented||$('ini'))return;var t=e.target,k=e.key;if(campo(t))return;
var ctl=/^(button|a|summary)$/i.test(t.tagName||'');
if(k=='l'||k=='L'){e.preventDefault();modo(!ler(),true);return}
if(k=='f'||k=='F'){fs();return}
if(k=='n'||k=='N'){if(!ler())nt();return}
if(ler())return;
if((k==' '||k=='Enter')&&ctl)return;
if(['ArrowRight','PageDown',' ','Enter'].indexOf(k)>-1){e.preventDefault();g(i+1)}
else if(['ArrowLeft','PageUp','Backspace'].indexOf(k)>-1){e.preventDefault();g(i-1)}
else if(k=='Home'){e.preventDefault();g(0)}else if(k=='End'){e.preventDefault();g(N-1)}});
function clique(id,fn){var b=$(id);if(b)b.addEventListener('click',fn)}
clique('p',function(){g(i-1)});clique('n',function(){g(i+1)});clique('b1',nt);clique('b2',fs);clique('bLer',function(){modo(!ler(),true)});clique('imp',function(){print()});
addEventListener('resize',function(){f();ajustar(S[i])});
D.addEventListener('conteudista:leitura',function(){ajustar(S[i])});
if(D.fonts&&D.fonts.ready)D.fonts.ready.then(function(){ajustar(S[i])});
var h=parseInt(location.hash.slice(1));i=isNaN(h)?0:Math.max(0,Math.min(N-1,h-1));modo(ler(),false);if(ler()&&i>0)S[i].scrollIntoView();
var x0=null;addEventListener('touchstart',function(e){x0=e.touches[0].clientX});addEventListener('touchend',function(e){if(x0===null||ler()){x0=null;return}var d=e.changedTouches[0].clientX-x0;if(Math.abs(d)>50)g(i+(d<0?1:-1));x0=null})})();"""
# Modo dos slides: telas pequenas abrem no modo leitura, salvo escolha anterior. Aplicado antes da pintura (sem piscar os slides).
SLIDES_MODO_HEAD='(function(){var m=null;try{m=localStorage.getItem("conteudista.apresentacao.modo")}catch(e){}if(m!=="leitura"&&m!=="slides")m=(window.matchMedia&&matchMedia("(max-width: 900px)").matches)?"leitura":"slides";if(m==="leitura")document.documentElement.classList.add("ml")})();'
SLIDES_MODO_BODY='if(document.documentElement.classList.contains("ml"))document.body.classList.add("modo-leitura");'
CSS_S_A11Y="""
.s{font-family:var(--leitura)}.slide-num{display:none}
.nt-t{font-weight:700;margin:0 0 4px!important;font-size:inherit!important}.nt-x{margin:0!important;font-size:inherit!important;line-height:1.45!important;white-space:pre-line}
.sn .s.on .nt{max-height:42%;overflow:auto;z-index:2}
.skip{position:absolute;left:8px;top:-80px;z-index:60;background:#fff;color:#0b3d6e;border:2px solid #0b3d6e;padding:8px 14px;border-radius:4px;font-weight:700;text-decoration:underline}
.skip:focus{top:8px}
#st:focus{outline:none}
#hud{flex-wrap:wrap;justify-content:flex-end;max-width:calc(100vw - 28px);z-index:10}
#hud label{color:#fff;font-weight:600}
#hud select{background:#fff;color:#0b3d6e;border:1px solid #fff;border-radius:6px;padding:5px 6px;font:inherit;font-size:13px}
#hud button:focus-visible,#hud select:focus-visible,#ini button:focus-visible{outline:3px solid #ffcd07;outline-offset:2px}
.skip:focus-visible,.s .in:focus-visible,.s .tbl:focus-visible,.s pre.code:focus-visible,.s summary:focus-visible{outline:3px solid #0b3d6e;outline-offset:2px}
.in.rola{overflow-y:auto}.s .tbl{overflow-x:auto}.s pre.code{overflow:auto;background:var(--soft);padding:10px;border-radius:8px;font-size:17px}
.s details.fluxo-tabela{font-size:16px;margin:6px 0;text-align:left}.s details.fluxo-tabela summary{cursor:pointer;color:var(--pri);font-weight:600}
.s details.fluxo-tabela table,.s details.fluxo-tabela p{font-size:16px}.fluxo-print{display:none}
@media (max-width:900px){#hud{left:8px;right:8px;justify-content:center;max-width:none}}
html.ml,html.ml body{overflow:auto!important;height:auto!important}
body.modo-leitura{background:#fff!important;color:#111!important;overflow:auto!important}
body.modo-leitura #st{position:static!important;transform:none!important;width:auto!important;height:auto!important;max-width:900px;margin:0 auto;padding:1rem 1rem 2rem}
body.modo-leitura .s{display:block!important;position:relative!important;inset:auto!important;width:auto!important;height:auto!important;min-height:0!important;overflow:visible!important;margin:0 0 2rem;padding:2rem 1.5rem;border:1px solid #bbb;border-radius:10px;background:#fff!important;color:#111!important;box-shadow:none!important;visibility:visible!important}
body.modo-leitura .s:before,body.modo-leitura .s:after{display:none!important}
body.modo-leitura .s.title,body.modo-leitura .s.highlight{background:#fff!important;color:#111!important}
body.modo-leitura .s.title h1,body.modo-leitura .s.title h2,body.modo-leitura .s.title p,body.modo-leitura .s.title strong{color:#111!important}
body.modo-leitura .s h1,body.modo-leitura .s h2,body.modo-leitura .s h3,body.modo-leitura .s strong{color:#0b3d6e!important}
body.modo-leitura .s h1{font-size:2rem!important;border-bottom:3px solid #0b3d6e!important;display:block!important}
body.modo-leitura .s h2{font-size:1.5rem!important}body.modo-leitura .s h3{font-size:1.2rem!important}
body.modo-leitura .s a{color:#0b3d6e!important;text-decoration:underline}
body.modo-leitura .s p,body.modo-leitura .s li{font-size:1.125rem!important;line-height:1.7!important;color:#111!important}
body.modo-leitura .s table{font-size:1rem!important;background:#fff!important}
body.modo-leitura .s th{background:#e8eef7!important;color:#111!important;border:1px solid #767676}
body.modo-leitura .s td{color:#111!important;border:1px solid #bbb}
body.modo-leitura .s code{background:#f4f4f4!important;color:#111!important}
body.modo-leitura .s blockquote{background:#f4f4f4!important;border-left-color:#555!important}
body.modo-leitura .s .in{position:static!important;display:block!important;height:auto!important;padding:0!important;margin:0!important;width:auto!important;overflow:visible!important}
body.modo-leitura .s .conteudo{transform:none!important;width:auto!important}
body.modo-leitura .s .nt{display:block!important;position:static!important;margin-top:1.25rem;padding:1rem 1.1rem;background:#f4f4f4!important;border:1px solid #bbb;border-left:5px solid #555;border-radius:6px;color:#111!important;font-size:1rem!important;line-height:1.55!important;max-height:none!important}
body.modo-leitura .slide-num{display:block!important;margin:0 0 .75rem;color:#555!important;font-size:.875rem;text-transform:uppercase;letter-spacing:.06em}
body.modo-leitura #hud{position:sticky!important;top:0;right:auto!important;bottom:auto!important;left:auto!important;z-index:20;max-width:900px;margin:0 auto;border-radius:0 0 10px 10px;background:#fff!important;color:#111!important;border:1px solid #bbb;border-top:0;opacity:1!important;flex-wrap:wrap;justify-content:center;padding:8px}
body.modo-leitura #hud label{color:#111}
body.modo-leitura #hud button{background:#fff!important;color:#0b3d6e!important;border:1px solid #0b3d6e!important}
body.modo-leitura #hud button[aria-pressed="true"]{background:#0b3d6e!important;color:#fff!important}
body.modo-leitura #hud select{background:#fff!important;color:#111!important;border:1px solid #555}
body.modo-leitura #hud button:focus-visible,body.modo-leitura #hud select:focus-visible{outline:3px solid #0b3d6e}
body.modo-leitura #bar,body.modo-leitura #p,body.modo-leitura #n,body.modo-leitura #c,body.modo-leitura #b1,
body.modo-leitura #bn,body.modo-leitura #ba,body.modo-leitura #ini{display:none!important}
@media print{
@page{size:A4;margin:18mm 16mm}
html,body,html.ml,html.ml body,body.modo-leitura{overflow:visible!important;height:auto!important;background:#fff!important}
#hud,#bar,#ini,.skip,.tem-som{display:none!important}
#st{position:static!important;transform:none!important;width:auto!important;height:auto!important;max-width:none!important;margin:0!important;padding:0!important}
#st .s{display:block!important;position:relative!important;inset:auto!important;width:auto!important;height:auto!important;min-height:0!important;overflow:visible!important;margin:0!important;padding:0 0 1.5rem!important;border:0!important;border-radius:0!important;background:#fff!important;color:#111!important;box-shadow:none!important;visibility:visible!important;break-after:page;page-break-after:always}
#st .s:last-child{break-after:auto;page-break-after:auto}
#st .s:before,#st .s:after{display:none!important}
#st .s .in{position:static!important;display:block!important;height:auto!important;padding:0!important;margin:0!important;width:auto!important;overflow:visible!important}
#st .s .conteudo{transform:none!important;width:auto!important}
#st .s h1,#st .s h2,#st .s h3,#st .s p,#st .s li,#st .s td,#st .s th,#st .s strong{color:#111!important}
#st .s h1{font-size:20pt!important;display:block!important}#st .s h2{font-size:15pt!important}
#st .s p,#st .s li{font-size:12pt!important;line-height:1.5!important}
#st .s th{background:#e8eef7!important}#st .s table{background:#fff!important}
#st .s h1,#st .s h2{border-color:#111!important}
#st .s a{color:#0b3d6e!important;text-decoration:underline!important}
#st .s .nt{display:block!important;position:static!important;margin-top:1rem!important;padding:.75rem 1rem!important;background:#f4f4f4!important;border:1px solid #bbb!important;border-left:5px solid #555!important;color:#111!important;font-size:10pt!important;max-height:none!important;break-inside:auto}
#st .s img{max-width:100%!important;max-height:none!important}
#st .s table{font-size:10pt!important}
#st .s figure,#st .s table,#st .s blockquote,#st .s pre{break-inside:avoid}
#st .s pre{white-space:pre-wrap!important;overflow-wrap:anywhere!important}#st .s .mw{overflow:visible!important}
#st .s .slide-num{display:block!important;color:#555!important;font-size:9pt}
details.fluxo-tabela{display:none!important}.fluxo-print{display:block!important}
}
"""
# Leitura acessível da apostila (1.66): vem depois do CSS do tema; não muda cores de identidade, só texto/foco/contraste.
CSS_G_A11Y="""
:root{--mu:#4f5b6a}[data-theme=dark]{--mu:#b8c4d2}[data-theme=contrast]{--mu:#fff}
.skip{position:absolute;left:8px;top:-80px;z-index:50;background:#fff;color:#0b3d6e;border:2px solid #0b3d6e;padding:8px 14px;border-radius:4px;font-weight:700;text-decoration:underline}
.skip:focus{top:8px}
header{flex-wrap:wrap}header label.lt{font-size:13px;color:inherit;font-weight:600}
@media (prefers-reduced-motion:reduce){html{scroll-behavior:auto}}
main:focus{outline:none}
article{font-family:var(--leitura);font-size:1.125rem;line-height:1.7;text-align:left;hyphens:none;-webkit-hyphens:none;overflow-wrap:break-word}
article p{margin:0 0 .85em}article li{margin:.3em 0}
article h1,article h2,article h3{font-family:var(--leitura);text-align:left;line-height:1.25}
article a{color:var(--pri);text-decoration:underline;text-underline-offset:.15em}
article code,article pre.code{font-family:Consolas,monospace}
a:focus-visible,button:focus-visible,select:focus-visible,summary:focus-visible,input:focus-visible,[tabindex]:focus-visible{outline:3px solid currentColor;outline-offset:3px}
article .tbl:focus-visible,article pre.code:focus-visible{outline:3px solid var(--pri);outline-offset:2px}
figcaption{color:var(--mu)}
.ab{background:var(--pri);color:var(--pa)}
details.fluxo-tabela{margin:-6px 0 16px}details.fluxo-tabela>summary{cursor:pointer;color:var(--pri);font-weight:600;padding:4px 0}
.fluxo-print{display:none}
li.task label{cursor:pointer}
@media print{
@page{size:A4;margin:18mm 16mm}
header,.skip,.ab{display:none!important}
.wrap{display:block!important}
nav{display:block!important;position:static!important;height:auto!important;overflow:visible!important;border:0!important;padding:0!important;font-size:11pt;color:#000!important}
nav a{color:#000!important;background:none!important;padding:2px 0!important}
main{break-before:page;padding:0!important}
body{background:#fff!important;color:#000!important}
article{border:0!important;box-shadow:none!important;padding:0!important;max-width:none!important;background:#fff!important;color:#000!important;font-family:var(--leitura);font-size:12pt;line-height:1.5}
article h1,article h2,article h3{color:#000!important}
article a{color:#000!important;text-decoration:underline}
article table,article th,article td,article blockquote,article code,article pre{background:#fff!important;color:#000!important;border-color:#000!important}
article th{font-weight:700}
article h2{break-before:page}
article h2,article h3{break-after:avoid}
article table,article figure,article blockquote,article pre,article .mw{break-inside:avoid}
article p{orphans:3;widows:3}
article pre.code{white-space:pre-wrap;overflow-wrap:break-word}
.tbl,article .mw{overflow:visible!important}
article pre.mermaid{white-space:pre-wrap;overflow-wrap:anywhere}
details.fluxo-tabela{display:none!important}
.fluxo-print{display:block!important}
}
"""
def title_of(md,d):
    m=re.search(r'^#\s+(.+)$',md,re.M);return plain(m.group(1)).strip() if m else d
def treinamento_html(md,bar,tema=None):
    # Os temas ANAC 20 anos alteram somente a apresentação; a apostila permanece no padrão gov.br.
    T=tema_de('govbr' if str(tema or '').startswith('anac20-') else tema)
    toc=[];b=html_of(blocks(md),toc)
    hs=list(re.finditer(r'<h2 id="[^"]+">(.*?)</h2>',b))
    for k,m in enumerate(hs):
        if 'gabarito' in m.group(1).lower():
            # o título fica visível; só o conteúdo (#gab) fica oculto, desfocado e fora da árvore de acessibilidade até ser mostrado
            e=hs[k+1].start() if k+1<len(hs) else len(b)
            b=b[:m.start()]+'<section class="gab-sec">'+m.group(0)+'<div id="gab" class="hid" aria-hidden="true" inert>'+b[m.end():e]+'</div></section>'+b[e:];break
    t=title_of(md,'Apostila');tc=''.join('<li class="lv%d"><a href="#%s">%s</a></li>'%(v,s,E(x)) for v,s,x in toc)
    return ('<!DOCTYPE html><html lang="pt-BR" data-theme="light" data-leitura="padrao"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>%s</title>'
    '<script>%s</script><style>%s%s%s</style></head><body><a class="skip" href="#conteudo">Ir para o conteúdo</a>'
    '<header><span class="t">%s</span><span id="pg" role="status" aria-live="polite"></span><select id="th" aria-label="Tema"><option value="light">Tema claro</option><option value="dark">Tema escuro</option><option value="contrast">Alto contraste</option></select>'
    '%s<button id="rs" type="button">Limpar checklists</button><button type="button" onclick="print()" title="Imprimir ou salvar em PDF em A4">Imprimir / PDF</button></header>'
    '<div class="wrap"><nav aria-label="Sumário"><ul>%s</ul></nav><main id="conteudo" tabindex="-1"><article>%s</article></main></div><script>%s%s</script>%s</body></html>'
    )%(E(t),LEITURA_HEAD,fontes_css(),tcss(CSS_G,T,T['extra_g']),CSS_G_A11Y,E(bar or t),LEITURA_SEL,tc,b,LEITURA_JS,JS_G.replace('__K__',slug(t)[:40]),MER)
def split_slides(md):
    t='Apresentação';fm=re.match(r'^---\n(.*?)\n---\n',md,re.S)
    if fm:
        x=re.search(r'^title:\s*(.+)$',fm.group(1),re.M);t=x.group(1).strip() if x else t;md=md[fm.end():]
    out=[]
    for p in [p.strip() for p in re.split(r'\n---\s*\n',md) if p.strip()]:
        c='';n=''
        m=re.search(r'<!--\s*class:\s*(\w+)\s*-->',p)
        if m:c=m.group(1);p=p.replace(m.group(0),'')
        m=re.search(r'<!--\s*notas:\s*(.*?)-->',p,re.S)
        if m:n=m.group(1).strip();p=p.replace(m.group(0),'')
        p=re.sub(r'<!--.*?-->','',p,flags=re.S).strip()
        if p:out.append((c,blocks(p),n))
    return t,out
def notas_para_narracao(md):
    """Texto a narrar em cada slide (mesma ordem de split_slides). <!-- narracao: texto --> substitui a nota;
    <!-- narracao: nao --> deixa o slide sem áudio. Slides sem notas: ''."""
    fm=re.match(r'^---\n(.*?)\n---\n',md or '',re.S)
    if fm:md=md[fm.end():]
    out=[]
    for p in [p.strip() for p in re.split(r'\n---\s*\n',md or '') if p.strip()]:
        n='';m=re.search(r'<!--\s*notas:\s*(.*?)-->',p,re.S)
        if m:n=m.group(1).strip()
        m=re.search(r'<!--\s*narra[cç][aã]o:\s*(.*?)-->',p,re.S|re.I)
        if m:
            v=m.group(1).strip();n='' if v.lower() in('nao','não','no','off','0','-') else v
        if re.sub(r'<!--.*?-->','',p,flags=re.S).strip():out.append(n)
    return out
def _narr(narr):
    """Normaliza a configuração de narração: {'html':{idx:caminho},'pptx':{idx:caminho},'avancar':bool,'duracoes':{idx:s}}."""
    if not narr:return {'html':{},'pptx':{},'avancar':False,'duracoes':{}}
    if isinstance(narr,str):
        import json as _j;narr=_j.loads(narr) if narr.strip() else {}
    f=lambda d:{int(k):v for k,v in (d or {}).items() if v and os.path.isfile(v)}
    dur={}
    for k,v in (narr.get('duracoes') or {}).items():
        try:dur[int(k)]=float(v)
        except Exception:pass
    return {'html':f(narr.get('html')),'pptx':f(narr.get('pptx')),'avancar':bool(narr.get('avancar')),'duracoes':dur}
MIME_AUDIO={'.mp3':'audio/mpeg','.ogg':'audio/ogg','.oga':'audio/ogg','.opus':'audio/ogg','.wav':'audio/wav','.m4a':'audio/mp4'}
CSS_NARR="""#ini{position:fixed;inset:0;z-index:20;background:rgba(7,29,65,.82);display:flex;align-items:center;justify-content:center;flex-direction:column;gap:14px;color:#fff;font-family:inherit}
#ini .it{font-size:22px;font-weight:700;margin:0}
#ini button{font-family:inherit;font-weight:600;font-size:22px;padding:14px 28px;border-radius:100em;border:0;background:#fff;color:#0b3d6e;cursor:pointer}
#ini button.ini0{background:transparent;color:#fff;font-size:16px;text-decoration:underline;padding:8px 16px}
#hud button[aria-pressed=true]{background:rgba(255,255,255,.4)}.s .tem-som{position:absolute;right:16px;top:12px;font-size:18px;opacity:.55}
@media print{#ini{display:none}.s .tem-som{display:none}}"""
JS_NARR="""(function(){var D=document,S=[].slice.call(D.querySelectorAll('.s')),A={},on=false,auto=__AUTO__,cur=-1,t=null,bn=D.getElementById('bn'),ba=D.getElementById('ba'),ini=D.getElementById('ini');
function ler(){return D.body.classList.contains('modo-leitura')}
function prox(){var n=D.getElementById('n');if(cur<S.length-1&&n&&!ler())n.click()}
[].slice.call(D.querySelectorAll('audio.narr')).forEach(function(a){A[+a.dataset.s]=a;a.addEventListener('ended',function(){if(auto&&on&&!ler()&&+a.dataset.s===cur){t=setTimeout(prox,700)}})});
function atual(){for(var k=0;k<S.length;k++)if(S[k].classList.contains('on'))return k;return 0}
function parar(){clearTimeout(t);Object.keys(A).forEach(function(x){A[x].pause()})}
function tocar(k){parar();Object.keys(A).forEach(function(x){A[x].currentTime=0});cur=k;if(!on||ler())return;if(A[k])A[k].play().catch(function(){});else if(auto)t=setTimeout(prox,4000)}
function marca(){if(bn){bn.setAttribute('aria-pressed',on?'true':'false');bn.textContent=on?'🔊 Narração (M)':'🔇 Narração (M)'}
if(ba){ba.setAttribute('aria-pressed',auto?'true':'false');ba.textContent=auto?'⏭ Avanço automático: sim (A)':'⏭ Avanço automático: não (A)'}}
function narr(v){on=v;marca();if(on)tocar(atual());else parar()}
D.addEventListener('conteudista:slide',function(e){tocar(e.detail)});
D.addEventListener('conteudista:modo',function(e){if(e.detail==='leitura')narr(false)});
if(bn)bn.addEventListener('click',function(){narr(!on)});
if(ba)ba.addEventListener('click',function(){auto=!auto;marca();if(auto&&on&&A[cur]&&A[cur].ended)tocar(cur)});
function fecharIni(som){if(!ini)return;var x=ini;ini=null;x.remove();narr(som);var st=D.getElementById('st');if(st)st.focus()}
if(ini){if(ler())fecharIni(false);else{var i1=D.getElementById('ini1'),i0=D.getElementById('ini0');
if(i1)i1.addEventListener('click',function(){fecharIni(true)});if(i0)i0.addEventListener('click',function(){fecharIni(false)});setTimeout(function(){if(ini&&i1)i1.focus()},30)}}
function campo(e){var n=(e.target.tagName||'').toLowerCase();return e.ctrlKey||e.metaKey||e.altKey||n=='input'||n=='select'||n=='textarea'||e.target.isContentEditable}
addEventListener('keydown',function(e){if(ini){var i1=D.getElementById('ini1'),i0=D.getElementById('ini0');
if(e.key==='Escape'){e.preventDefault();fecharIni(false)}else if(e.key==='Tab'){e.preventDefault();(D.activeElement===i1?i0:i1).focus()}return}
if(campo(e))return;if(e.key=='m'||e.key=='M'){if(!ler())narr(!on)}else if((e.key=='a'||e.key=='A')&&ba){ba.click()}},true);
S.forEach(function(x,k){if(A[k]){var m=D.createElement('span');m.className='tem-som';m.title='Este slide tem narração';m.setAttribute('aria-hidden','true');m.textContent='🔊';x.appendChild(m)}});marca()})();"""
def _anac_html_css(tema,n):
    if not str(tema or '').startswith('anac20-'):return ''
    pasta=os.path.join(os.path.dirname(os.path.abspath(__file__)),'temas',tema)
    def data(nome):
        cam=os.path.join(pasta,nome)
        return 'data:image/png;base64,'+base64.b64encode(open(cam,'rb').read()).decode('ascii') if os.path.isfile(cam) else ''
    imgs=[data('capa.png'),data('conteudo-a.png'),data('conteudo-b.png'),data('encerramento.png')]
    escuro=tema=='anac20-escuro'; tinta='#ffffff' if escuro else '#193f52'; mut='#dce8ee' if escuro else '#526d7c'
    regras=['.s{background-size:cover!important;background-position:center!important;background-repeat:no-repeat!important;color:%s!important}'%tinta,
      '.s h1,.s h2,.s h3,.s p,.s li{color:%s!important}.s .in{position:relative;z-index:1}.s:not(.title) .in{padding-top:1.2vh}'%tinta,
      '.s table{background:rgba(255,255,255,.94);color:#193f52}.s table *{color:#193f52!important}.s .nt{color:%s!important}'%mut]
    for i in range(n):
        fundo=imgs[0] if i==0 else imgs[3] if i==n-1 and 'title' else imgs[1 if i%2 else 2]
        if fundo:regras.append('.s:nth-child(%d){background-image:url("%s")!important}'%(i+1,fundo))
    if escuro:regras.append('.s.title .in{margin-left:37%;width:58%}.s.title h1{color:#fff!important}')
    else:regras.append('.s.title .in{margin-left:37%;width:58%}.s.title h1{color:#193f52!important}')
    return '\n'.join(regras)
MER_S='<script type="module">try{const m=(await import("https://cdn.jsdelivr.net/npm/mermaid@10/dist/mermaid.esm.min.mjs")).default;m.initialize({startOnLoad:false});document.body.classList.add("mm");try{await m.run({querySelector:"pre.mermaid:not([data-processed])",suppressErrors:true})}finally{document.body.classList.remove("mm")}}catch(e){document.body.classList.remove("mm");console.warn("Mermaid offline",e)}</script>'
def _mer_s(o):
    """Mermaid dos slides: so e carregado se algum fluxo ficou como codigo (sem imagem). Desenha com todos os slides mensuraveis (classe mm no CSS): o Mermaid nao desenha dentro de elemento oculto."""
    return MER_S if 'class="mermaid"' in o else ''
def _fluxos_dict(f):
    """Aceita o dicionario {codigo do fluxo: imagem em data URI} ou o mesmo em JSON (texto)."""
    if isinstance(f,str):
        try:
            import json;f=json.loads(f)
        except Exception:return {}
    return {str(k).strip():v for k,v in f.items() if isinstance(v,str) and v.startswith('data:image/')} if isinstance(f,dict) else {}
def _fluxo_alt(cod):
    """Texto alternativo do fluxo: as etapas na ordem em que aparecem."""
    try:
        vistos=[]
        for r in mermaid_edges(cod)[1:]:
            for x in (r[0],r[2]):
                if x and x not in vistos:vistos.append(x)
        return ('Fluxo: '+'; '.join(vistos))[:700] if vistos else 'Fluxo'
    except Exception:return 'Fluxo'
def _fluxos_img(o,fluxos):
    """Troca cada fluxo Mermaid do HTML dos slides pela imagem pre-renderizada do mesmo codigo, quando ela existe (a tabela do fluxo continua)."""
    if not fluxos:return o
    def rep(m):
        cod=html.unescape(m.group(1)).strip();img=fluxos.get(cod)
        if not img:return m.group(0)
        return '<figure class="mw fx fluxo"><img class="fxi" alt="%s" src="%s"><figcaption>Fluxo</figcaption></figure>'%(E(_fluxo_alt(cod)),img)
    return re.sub(r'<div class="mw"[^>]*><pre class="mermaid">(.*?)</pre></div>',rep,o,flags=re.S)
def _slides_doc(t,sl,foot,tema=None,audios=None,avancar=False,fluxos=None):
    """Monta o HTML da apresentação (com ou sem narração): ARIA, leitura, modo leitura, auto-ajuste e impressão A4."""
    T=tema_de(tema);n=len(sl);secs=[]
    for k,(c,b,nt) in enumerate(sl):
        rot='Slide %d de %d'%(k+1,n)
        nota='<aside class="nt" role="note" aria-label="Notas do apresentador"><p class="nt-t">Notas do apresentador</p><p class="nt-x">%s</p></aside>'%E(nt) if (nt or '').strip() else ''
        secs.append('<section class="s %s" role="group" aria-roledescription="slide" aria-label="%s" tabindex="-1"><p class="slide-num" aria-hidden="true">%s</p><div class="in"><div class="conteudo">%s</div></div>%s</section>'%(c,rot,rot,html_of(b),nota))
    o=_fluxos_img(''.join(secs),fluxos)
    css=fontes_css()+tcss(CSS_S.replace('__F__',foot.replace('"','')),T,T['extra_s'])+_anac_html_css(tema,n)
    narr=bool(audios);aud=''
    if narr:
        css+=CSS_NARR
        for k,cam in sorted(audios.items()):
            if 0<=k<n:
                mime=MIME_AUDIO.get(os.path.splitext(cam)[1].lower(),'audio/mpeg')
                aud+='<audio class="narr" data-s="%d" preload="auto" src="data:%s;base64,%s"></audio>'%(k,mime,base64.b64encode(open(cam,'rb').read()).decode())
    css+=CSS_S_A11Y
    hud=('<div id="hud" role="toolbar" aria-label="Controles da apresentação"><button id="p" type="button" aria-label="Slide anterior">◀</button>'
         '<span id="c" role="status" aria-live="polite" aria-atomic="true"></span><button id="n" type="button" aria-label="Próximo slide">▶</button>'
         +('<button id="bn" type="button" aria-pressed="false">🔇 Narração (M)</button><button id="ba" type="button" aria-pressed="false">⏭ Avanço automático (A)</button>' if narr else '')
         +LEITURA_SEL+'<button id="b1" type="button" aria-pressed="false">Notas (N)</button><button id="bLer" type="button" aria-pressed="false">Modo leitura</button>'
         '<button id="imp" type="button" title="Imprimir ou salvar em PDF em A4">Imprimir / PDF</button><button id="b2" type="button" aria-pressed="false">Tela cheia (F)</button></div>')
    ini=('<div id="ini" role="dialog" aria-modal="true" aria-labelledby="iniT"><p id="iniT" class="it">Iniciar a apresentação</p>'
         '<button id="ini1" type="button">▶ Iniciar com narração</button><button id="ini0" type="button" class="ini0">Iniciar sem som</button></div>') if narr else ''
    js='<script>'+LEITURA_JS+JS_S+'</script>'+('<script>'+JS_NARR.replace('__AUTO__','true' if avancar else 'false')+'</script>' if narr else '')
    return ('<!DOCTYPE html><html lang="pt-BR" data-leitura="padrao"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="theme" content="%s"><title>%s</title>'
    '<script>%s%s</script><style>%s</style></head><body><script>%s</script><a class="skip" href="#st">Ir para os slides</a>%s'
    '<div id="st" role="region" aria-roledescription="apresentação" aria-label="Apresentação" tabindex="-1">%s</div><div id="bar" aria-hidden="true"></div>%s%s%s%s</body></html>'
    )%(E(tema or TEMA_PADRAO),E(t),LEITURA_HEAD,SLIDES_MODO_HEAD,css,SLIDES_MODO_BODY,hud,o,ini,aud,js,_mer_s(o))
def slides_html(t,sl,foot,tema=None,audios=None,avancar=False,fluxos=None):
    return _slides_doc(t,sl,foot,tema,audios,avancar,fluxos)
def _slides_html_narrado(t,sl,foot,tema,audios,avancar,fluxos=None):
    """Compatibilidade: a montagem é a mesma de slides_html (_slides_doc); a narração escuta o evento conteudista:slide."""
    return _slides_doc(t,sl,foot,tema,audios,avancar,fluxos)
# ============================== DOCX ==============================
def treinamento_docx(md,path,foot,tema=None):
    # Os temas ANAC 20 anos alteram somente a apresentação; o DOCX permanece no padrão gov.br.
    # Leitura acessível (1.66): o DOCX é único, em Verdana 12 pt, entrelinha 1,5 e 6 pt após; T['docx_font'] não escolhe mais a fonte.
    T=tema_de('govbr' if str(tema or '').startswith('anac20-') else tema);C=lambda h:trgb(h,T);H=lambda h:thex(h,T);FONTE='Verdana'
    from docx import Document
    from docx.shared import Pt,RGBColor,Cm
    from docx.enum.text import WD_BREAK
    from docx.oxml.ns import qn
    def novo_num(doc):
        "Cria uma numeracao nova (reiniciada em 1) para cada lista numerada; sem isso o Word continua a contagem da lista anterior."
        sn=doc.styles['List Number'].element.pPr.numPr.numId.val
        nums=doc.part.numbering_part.numbering_definitions._numbering
        ab=[int(n.find(qn('w:abstractNumId')).get(qn('w:val'))) for n in nums.findall(qn('w:num')) if int(n.get(qn('w:numId')))==sn][0]
        n=nums.add_num(ab);n.add_lvlOverride(ilvl=0).add_startOverride(1);return n.numId
    def aplica_num(p,nid):
        np_=p._p.get_or_add_pPr().get_or_add_numPr();np_.get_or_add_ilvl().val=0;np_.get_or_add_numId().val=nid
    from docx.oxml import OxmlElement
    PRI=C('0B3D6E');ACC=C('0A7CC1')
    d=Document();sec=d.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21)
    for a in('left_margin','right_margin'):setattr(sec,a,Cm(2.2))
    sec.top_margin=Cm(2);sec.bottom_margin=Cm(2)
    stl=d.styles['Normal'];stl.font.name=FONTE;stl.font.size=Pt(12)
    pf=stl.paragraph_format;pf.line_spacing=1.5;pf.space_before=Pt(0);pf.space_after=Pt(6);pf.widow_control=True
    stl.element.rPr.rFonts.set(qn('w:eastAsia'),FONTE);stl.element.rPr.rFonts.set(qn('w:cs'),FONTE)
    def lang_pt(rpr):
        for x in rpr.findall(qn('w:lang')):rpr.remove(x)
        lg=OxmlElement('w:lang');lg.set(qn('w:val'),'pt-BR');lg.set(qn('w:eastAsia'),'pt-BR');lg.set(qn('w:bidi'),'pt-BR');rpr.append(lg)
    lang_pt(stl.element.get_or_add_rPr());d.core_properties.language='pt-BR'
    for nm,sz in(('Title',24),('Heading 1',16),('Heading 2',13),('Heading 3',12)):
        s=d.styles[nm];s.font.name=FONTE;s.font.size=Pt(sz);s.font.color.rgb=PRI;s.font.bold=True
    if True:  # força Verdana também nos títulos e listas (remove a fonte do tema do Word)
        for nm in('Title','Heading 1','Heading 2','Heading 3','List Bullet','List Number'):
            try:
                st_=d.styles[nm];st_.font.name=FONTE;rf=st_.element.get_or_add_rPr().get_or_add_rFonts()
                for at in('w:asciiTheme','w:hAnsiTheme','w:eastAsiaTheme','w:cstheme'):rf.attrib.pop(qn(at),None)
                rf.set(qn('w:eastAsia'),FONTE);rf.set(qn('w:cs'),FONTE)
            except KeyError:pass
        rf=stl.element.rPr.rFonts;rf.set(qn('w:eastAsia'),FONTE)
    def shade(el,fill):
        pr=el.get_or_add_tcPr() if hasattr(el,'get_or_add_tcPr') else el.get_or_add_pPr()
        sh=OxmlElement('w:shd');sh.set(qn('w:val'),'clear');sh.set(qn('w:color'),'auto');sh.set(qn('w:fill'),fill);pr.append(sh)
    def lborder(p,color):
        pPr=p._p.get_or_add_pPr();b=OxmlElement('w:pBdr');l=OxmlElement('w:left')
        for k,v in(('val','single'),('sz','24'),('space','8'),('color',color)):l.set(qn('w:'+k),v)
        b.append(l);pPr.append(b)
    def link(p,text,url,bold=False,size=None):
        r_id=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
        h=OxmlElement('w:hyperlink');h.set(qn('r:id'),r_id);r=OxmlElement('w:r');rp=OxmlElement('w:rPr')
        c=OxmlElement('w:color');c.set(qn('w:val'),H('0A7CC1'));u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(c);rp.append(u)
        if bold:rp.append(OxmlElement('w:b'))
        if size:
            z=OxmlElement('w:sz');z.set(qn('w:val'),str(int(size*2)));rp.append(z)
        r.append(rp);t=OxmlElement('w:t');t.text=text;t.set(qn('xml:space'),'preserve');r.append(t);h.append(r);p._p.append(h)
    def pic(run,src,maxcm=16.0,hcm=None,alt=''):
        raw=img_bytes(src)
        if raw is None:run.text='[imagem%s]'%(': '+alt if alt else '');return False
        w,h=img_size(raw)
        try:
            if hcm:shp=run.add_picture(io.BytesIO(raw),height=Cm(hcm))
            else:shp=run.add_picture(io.BytesIO(raw),width=Cm(min(maxcm,w/96*2.54)))
            # texto alternativo: vem do Markdown; vazio continua vazio (não se inventa descrição) e gera aviso
            dp=shp._inline.docPr;dp.set('descr',alt or '');dp.set('title',alt or '')
            if not alt:aviso('Imagem sem descrição (texto alternativo vazio) na apostila DOCX.')
            return True
        except Exception:run.text='[imagem]';return False
    def fill(p,t,color=None,size=None,bold=False):
        for x,b,i,c,u in runs(t):
            if c==2:pic(p.add_run(),u,hcm=0.5,alt=x);continue
            if u:link(p,x,u,b,size);continue
            r=p.add_run(x);r.bold=bool(b) or bold;r.italic=bool(i)
            if c:r.font.name='Consolas';r.font.size=Pt(10);r.font.color.rgb=PRI
            if color is not None:r.font.color.rgb=color
            if size:r.font.size=Pt(size)
    def table(rows):
        nc=max(len(r) for r in rows);tb=d.add_table(rows=len(rows),cols=nc);tb.style='Table Grid'
        for ri,r in enumerate(rows):
            for ci in range(nc):
                cell=tb.cell(ri,ci);cell.text='';p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.15
                fill(p,r[ci] if ci<len(r) else '',size=10.5,bold=(ri==0))
                if ri==0:
                    shade(cell._tc,H('EAF1F8'))
                    for x in p.runs:x.font.color.rgb=PRI
        tr=tb.rows[0]._tr;trPr=tr.get_or_add_trPr();h=OxmlElement('w:tblHeader');h.set(qn('w:val'),'true');trPr.append(h)
        d.add_paragraph().paragraph_format.space_after=Pt(2)
    def render(bl,inq=None):
        for b in bl:
            k=b[0]
            if k=='h':
                if b[1]==1:
                    p=d.add_paragraph(style='Title');fill(p,b[2],color=PRI);continue
                if b[1]==2 and 'gabarito' in b[2].lower():d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                p=d.add_heading(level=min(b[1]-1,3));fill(p,b[2])
                for r in p.runs:r.font.color.rgb=PRI if b[1]==2 else ACC
            elif k=='p':
                p=d.add_paragraph();fill(p,b[1])
                if inq:shade(p._p,inq[0]);lborder(p,inq[1])
            elif k=='img':
                p=d.add_paragraph();p.alignment=1;pic(p.add_run(),b[2],alt=b[1])
                if b[1]:
                    c=d.add_paragraph();c.alignment=1;r=c.add_run(b[1]);r.italic=True;r.font.size=Pt(9)
            elif k=='table':table(b[1])
            elif k=='list':
                nid=None
                for dp,od,ck,t in b[1]:
                    if ck is not None:
                        p=d.add_paragraph();p.paragraph_format.left_indent=Cm(0.6+0.6*dp);p.add_run(('☒ ' if ck else '☐ ')).font.size=Pt(12);fill(p,t)
                    else:
                        st='List Number' if od else 'List Bullet'
                        if dp:st+=' %d'%(min(dp,2)+1)
                        p=d.add_paragraph(style=st);fill(p,t)
                        if od and not dp:
                            if nid is None:nid=novo_num(d)
                            aplica_num(p,nid)
                    p.paragraph_format.space_after=Pt(6)
                    if inq:shade(p._p,inq[0]);lborder(p,inq[1])
            elif k=='quote':
                col={'tip':(H('E9F6F0'),H('1E7B4F')),'warning':(H('FFF4E0'),H('9A5B00'))}.get(b[1],(H('EAF1F8'),H('0A7CC1')))
                render(b[2],col)
            elif k=='code':
                if b[1]=='mermaid':
                    p=d.add_paragraph();r=p.add_run('Fluxo (representação em tabela):');r.bold=True;r.font.color.rgb=ACC
                    table(mermaid_edges(b[2]))
                else:
                    p=d.add_paragraph();p.paragraph_format.line_spacing=1.15;r=p.add_run(b[2]);r.font.name='Consolas';r.font.size=Pt(9.5);shade(p._p,H('F1F4F8'))
    bl=blocks(sem_comentarios(md))
    hi=next((i for i,x in enumerate(bl) if x[0]=='h' and x[1]==2),len(bl))
    render(bl[:hi])
    p=d.add_paragraph();r=p.add_run('Sumário');r.bold=True;r.font.size=Pt(14);r.font.color.rgb=PRI
    p=d.add_paragraph();r=p.add_run()
    for tag,txt in(('begin',None),(None,'TOC \\o "1-2" \\h \\z \\u'),('separate',None),(None,None),('end',None)):
        if tag:
            f=OxmlElement('w:fldChar');f.set(qn('w:fldCharType'),tag);r._r.append(f)
        elif txt:
            it=OxmlElement('w:instrText');it.set(qn('xml:space'),'preserve');it.text=txt;r._r.append(it)
        else:
            t=OxmlElement('w:t');t.text='Clique com o botão direito e escolha "Atualizar campo" para gerar o sumário.';r._r.append(t)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    render(bl[hi:])
    fp=sec.footer.paragraphs[0];fp.text=foot or ' ';fp.runs[0].font.size=Pt(8);fp.runs[0].font.color.rgb=C('5A6676')
    s=d.settings.element;u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');s.append(u)
    for r_ in list(d.element.body.iter(qn('w:r')))+list(sec.footer._element.iter(qn('w:r'))):  # idioma pt-BR em todos os trechos (texto, links, tabelas, legendas)
        rpr=r_.find(qn('w:rPr'))
        if rpr is None:rpr=OxmlElement('w:rPr');r_.insert(0,rpr)
        lang_pt(rpr)
    d.save(path)
# ============================== EPUB (apostila) ==============================
EPUB_CSS="""@font-face{font-family:"__NOME__";src:url("../fonts/__BASE__-400.woff2") format("woff2");font-weight:400;font-style:normal;font-display:swap}
html{font-size:100%}
body{font-family:"__NOME__",Verdana,Tahoma,Arial,sans-serif;font-size:1em;line-height:1.65;color:#111;background:#fff;margin:0;padding:0;text-align:left;hyphens:none;-webkit-hyphens:none;overflow-wrap:break-word}
p{margin:0 0 1em}
h1,h2,h3,h4{color:#0b3d6e;line-height:1.25;text-align:left;break-after:avoid;page-break-after:avoid}
h1{font-size:1.8em}h2{font-size:1.45em;margin:1.4em 0 .6em}h3{font-size:1.15em;margin:1.2em 0 .4em}h4{font-size:1.05em}
a{color:#0b3d6e;text-decoration:underline}
table{width:100%;border-collapse:collapse;margin:1em 0;font-size:.95em;table-layout:fixed}
th,td{border:1px solid #767676;padding:.4em .5em;text-align:left;vertical-align:top;word-break:break-word}
th{background:#e8eef7;color:#111}
figure{margin:1em 0;text-align:center}
img{max-width:100%;height:auto}
img.il{max-height:1.4em;vertical-align:middle}
figcaption{font-size:.9em;color:#333;margin-top:.25em}
blockquote{margin:1em 0;padding:.75em 1em;border-left:4px solid #555;background:#f4f4f4}
pre{white-space:pre-wrap;overflow-wrap:break-word;background:#f4f4f4;padding:.75em;border:1px solid #ccc}
code{font-family:Consolas,monospace;font-size:.9em}
li.task{list-style:none;margin-left:-1.2em}
.caixa{font-size:1.2em;margin-right:.35em}
.sr-only{position:absolute;width:1px;height:1px;padding:0;margin:-1px;overflow:hidden;clip:rect(0,0,0,0);white-space:nowrap;border:0}
.fluxo-t{font-weight:700;margin-bottom:.25em}
.rodape{font-size:.9em;color:#333;border-top:1px solid #999;padding-top:.5em;margin-top:2em}
nav ol{list-style:none;padding-left:1em}
table,figure,blockquote,pre{break-inside:avoid;page-break-inside:avoid}
"""
def _ep_img_raw(src):
    "Bytes e media-type de uma imagem (data URI ou arquivo local). Imagem remota não entra no EPUB."
    try:
        if src.startswith('data:'):
            h,d_=src.split(',',1);mime=h[5:].split(';')[0].lower()
            raw=base64.b64decode(d_) if ';base64' in h else unquote_to_bytes(d_)
        elif os.path.isfile(src):raw=open(src,'rb').read();mime=(mimetypes.guess_type(src)[0] or '').lower()
        else:return None,None
    except Exception:return None,None
    if mime=='image/jpg':mime='image/jpeg'
    if mime not in('image/png','image/jpeg','image/gif','image/svg+xml','image/webp'):
        try:
            from PIL import Image
            b=io.BytesIO();Image.open(io.BytesIO(raw)).save(b,'PNG');raw=b.getvalue();mime='image/png'
        except Exception:return None,None
    return raw,mime
class _EpubCtx:
    def __init__(self):
        self.imgs={};self.ids={};self.sem_alt=0;self.n_img=0;self.sec_ids=[]
    def uid(self,base):
        base=base or 'secao';n=self.ids.get(base,0)+1;self.ids[base]=n
        return base if n==1 else '%s-%d'%(base,n)
    def img(self,src):
        raw,mime=_ep_img_raw(src)
        if raw is None:return None
        h=hashlib.sha1(raw).hexdigest()
        if h not in self.imgs:
            ext={'image/png':'png','image/jpeg':'jpg','image/gif':'gif','image/svg+xml':'svg','image/webp':'webp'}[mime]
            self.imgs[h]=('img/img-%03d.%s'%(len(self.imgs)+1,ext),mime,raw)
        return self.imgs[h][0]
def _ep_inl(t,cx):
    s=''
    for x,b,i,c,u in runs(t):
        if c==2:
            src=cx.img(u);cx.n_img+=1
            if not x.strip():cx.sem_alt+=1
            s+='<img class="il" src="%s" alt="%s"/>'%(E(src),E(x)) if src else E('[imagem%s]'%(': '+x if x else ''),quote=False)
            continue
        h=E(x,quote=False)
        if c:h='<code>'+h+'</code>'
        if i:h='<em>'+h+'</em>'
        if b:h='<strong>'+h+'</strong>'
        if u:h='<a href="%s">%s</a>'%(E(u),h)
        s+=h
    return s
def _ep_list(it,cx):
    def rec(k,d):
        tag='ol' if it[k][1] else 'ul';o=''
        while k<len(it) and it[k][0]>=d:
            if it[k][0]>d:sub,k=rec(k,it[k][0]);o=o[:-5]+sub+'</li>';continue
            dp,od,ck,t=it[k]
            if ck is None:o+='<li>'+_ep_inl(t,cx)+'</li>'
            else:o+='<li class="task"><span class="caixa" aria-hidden="true">%s</span><span class="sr-only">%s</span>%s</li>'%('☒' if ck else '☐','Concluído: ' if ck else 'Não concluído: ',_ep_inl(t,cx))
            k+=1
        return '<%s>%s</%s>'%(tag,o,tag),k
    return rec(0,it[0][0])[0]
def _ep_tabela(rows,cx,cap=''):
    nc=max(len(r) for r in rows)
    t='<table>'+('<caption>%s</caption>'%E(cap) if cap else '')+'<thead><tr>'+''.join('<th scope="col">%s</th>'%_ep_inl(rows[0][c] if c<len(rows[0]) else '',cx) for c in range(nc))+'</tr></thead><tbody>'
    t+=''.join('<tr>'+''.join('<td>%s</td>'%_ep_inl(r[c] if c<len(r) else '',cx) for c in range(nc))+'</tr>' for r in rows[1:])
    return t+'</tbody></table>'
def _ep_blocks(bl,cx,toc=None):
    o=[]
    for b in bl:
        k=b[0]
        if k=='h':
            sid=cx.uid(slug(b[2]));lv=min(b[1],6)
            if toc is not None and lv==3:toc.append((sid,plain(b[2])))
            o.append('<h%d id="%s">%s</h%d>'%(lv,sid,_ep_inl(b[2],cx),lv))
        elif k=='p':o.append('<p>'+_ep_inl(b[1],cx)+'</p>')
        elif k=='img':
            src=cx.img(b[2]);cx.n_img+=1
            if not b[1].strip():cx.sem_alt+=1
            if src:o.append('<figure><img src="%s" alt="%s"/>%s</figure>'%(E(src),E(b[1]),'<figcaption>%s</figcaption>'%E(b[1]) if b[1] else ''))
            else:o.append('<p>%s</p>'%E('[Imagem não incluída%s]'%(': '+b[1] if b[1] else ''),quote=False));aviso('Imagem remota ou ilegível não incluída no EPUB: %s'%b[2][:60])
        elif k=='table':o.append(_ep_tabela(b[1],cx))
        elif k=='list':o.append(_ep_list(b[1],cx))
        elif k=='quote':o.append('<blockquote class="%s">%s</blockquote>'%(E(b[1] or ''),_ep_blocks(b[2],cx,toc)))
        elif k=='code':
            if b[1]=='mermaid':
                try:rows=mermaid_edges(b[2])
                except Exception:rows=None
                if rows and len(rows)>1:o.append('<div class="fluxo"><p class="fluxo-t">Fluxo</p>%s</div>'%_ep_tabela(rows,cx))
                else:o.append('<pre><code>%s</code></pre>'%E(b[2],quote=False))
            else:o.append('<pre><code>%s</code></pre>'%E(b[2],quote=False))
    return '\n'.join(o)
def _ep_doc(title,body,tipo):
    return ('<?xml version="1.0" encoding="utf-8"?>\n<!DOCTYPE html>\n<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="pt-BR" xml:lang="pt-BR">\n'
            '<head>\n<meta charset="utf-8"/>\n<title>%s</title>\n<link rel="stylesheet" type="text/css" href="css/style.css"/>\n</head>\n<body epub:type="%s">\n%s\n</body>\n</html>\n')%(E(title),tipo,body)
def treinamento_epub(md,path,foot,tema=None,fonte='atkinson'):
    """Apostila em EPUB 3 (XHTML + OPF + nav + NCX), com a fonte de leitura embutida e metadados de acessibilidade.
    O EPUB é claro e neutro (não usa os temas escuros); o tema não altera o EPUB."""
    from lxml import etree
    nome_f,base_f,ofl=LEITURAS[fonte]
    md=sem_comentarios(md);bl=blocks(md);titulo=title_of(md,'Apostila');cx=_EpubCtx()
    hi=next((i for i,x in enumerate(bl) if x[0]=='h' and x[1]==2),len(bl))
    capa=bl[:hi];secoes=[];atual=None
    for x in bl[hi:]:
        if x[0]=='h' and x[1]==2:atual=[x];secoes.append(atual)
        else:atual.append(x)
    if not secoes:  # sem h2: a capa leva só o título; o corpo inteiro vai para um único conteudo.xhtml
        h1i=next((i for i,x in enumerate(bl) if x[0]=='h' and x[1]==1),-1)
        capa=bl[:h1i+1];resto=bl[h1i+1:]
        if resto:secoes=[resto]
    arquivos=[];toc=[];ids_arq={}
    # capa: título (h1), barra/subtítulo, parágrafos antes do primeiro h2 e rodapé
    corpo_capa=_ep_blocks(capa,cx)
    if '<h1' not in corpo_capa:corpo_capa='<h1 id="%s">%s</h1>\n'%(cx.uid('titulo'),E(titulo))+corpo_capa
    if foot:corpo_capa+='\n<p class="rodape">%s</p>'%E(foot)
    arquivos.append(('capa','capa.xhtml',titulo,'<section epub:type="titlepage">\n%s\n</section>'%corpo_capa,'frontmatter'))
    for n,sec_ in enumerate(secoes,1):
        nome='sec-%02d.xhtml'%n if secoes and sec_[0][0]=='h' and sec_[0][1]==2 else 'conteudo.xhtml'
        h3=[];corpo=_ep_blocks(sec_,cx,h3)
        m=re.search(r'<h2 id="([^"]+)">',corpo);sid=m.group(1) if m else ''
        tit=plain(sec_[0][2]) if sec_[0][0]=='h' else titulo
        toc.append((nome,sid,tit,h3));arquivos.append(('s%02d'%n,nome,tit,'<section epub:type="chapter">\n%s\n</section>'%corpo,'bodymatter'))
    # links internos (#id) apontam para o arquivo certo
    for _,nome,_,corpo,_ in arquivos:
        for i_ in re.findall(r'\bid="([^"]+)"',corpo):ids_arq.setdefault(i_,nome)
    def fixa(nome,corpo):
        return re.sub(r'href="#([^"]+)"',lambda m_:'href="%s#%s"'%(ids_arq[m_.group(1)],m_.group(1)) if ids_arq.get(m_.group(1),nome)!=nome else m_.group(0),corpo)
    docs={nome:_ep_doc(tit,fixa(nome,corpo),tipo) for _,nome,tit,corpo,tipo in arquivos}
    # sumário (nav) e landmarks
    li=['<li><a href="capa.xhtml">%s</a></li>'%E(titulo)]
    for nome,sid,tit,h3 in toc:
        sub='<ol>%s</ol>'%''.join('<li><a href="%s#%s">%s</a></li>'%(nome,i_,E(t_)) for i_,t_ in h3) if h3 else ''
        li.append('<li><a href="%s%s">%s</a>%s</li>'%(nome,'#'+sid if sid else '',E(tit),sub))
    inicio=toc[0][0] if toc else 'capa.xhtml'
    nav_body=('<nav epub:type="toc" id="toc" role="doc-toc">\n<h1>Sumário</h1>\n<ol>\n%s\n</ol>\n</nav>\n'
              '<nav epub:type="landmarks" id="landmarks" hidden="hidden">\n<h2>Pontos de referência</h2>\n<ol>\n'
              '<li><a epub:type="titlepage" href="capa.xhtml">Capa</a></li>\n<li><a epub:type="toc" href="nav.xhtml">Sumário</a></li>\n'
              '<li><a epub:type="bodymatter" href="%s">Início do conteúdo</a></li>\n</ol>\n</nav>')%('\n'.join(li),inicio)
    docs['nav.xhtml']=_ep_doc('Sumário',nav_body,'frontmatter')
    for nome_,x in docs.items():etree.fromstring(x.encode('utf-8'))  # XHTML bem formado (erro interrompe só este EPUB)
    # fontes da família escolhida (somente ela) + licença
    faces=fontes_faces(fonte);css=EPUB_CSS.split('\n',1)[1].replace('__NOME__',nome_f)
    ff=''.join('@font-face{font-family:"%s";src:url("../fonts/%s") format("woff2");font-weight:%d;font-style:%s;font-display:swap}\n'%(nome_f,arq,peso,estilo) for peso,estilo,arq,_ in faces)
    css=ff+css
    ident='urn:uuid:%s'%uuid.uuid5(uuid.NAMESPACE_URL,'conteudista:apostila:%s:%s'%(titulo,fonte))
    agora=datetime.datetime.now(datetime.timezone.utc)
    titulo_opf=titulo+(' — OpenDyslexic' if fonte=='opendyslexic' else '')
    man=['<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>','<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>',
         '<item id="css" href="css/style.css" media-type="text/css"/>']
    for k_,(peso,estilo,arq,_) in enumerate(faces,1):man.append('<item id="fonte%d" href="fonts/%s" media-type="font/woff2"/>'%(k_,arq))
    if faces and os.path.isfile(os.path.join(FONTES_DIR,ofl)):man.append('<item id="licenca" href="fonts/%s" media-type="text/plain"/>'%ofl)
    for k_,(nm,mime,_) in enumerate(cx.imgs.values(),1):man.append('<item id="img%d" href="%s" media-type="%s"/>'%(k_,nm,mime))
    for idr,nome,_,_,_ in arquivos:man.append('<item id="%s" href="%s" media-type="application/xhtml+xml"/>'%(idr,nome))
    spine=['<itemref idref="capa"/>','<itemref idref="nav"/>']+['<itemref idref="%s"/>'%idr for idr,nome,_,_,_ in arquivos[1:]]
    tem_img=cx.n_img>0;alt_ok=tem_img and cx.sem_alt==0
    if cx.sem_alt:aviso('EPUB: %d imagem(ns) sem texto alternativo; o recurso alternativeText não foi declarado.'%cx.sem_alt)
    meta=['<dc:identifier id="pub-id">%s</dc:identifier>'%ident,'<dc:title>%s</dc:title>'%E(titulo_opf),'<dc:language>pt-BR</dc:language>','<dc:creator>Conteudista</dc:creator>',
          '<dc:date>%s</dc:date>'%agora.strftime('%Y-%m-%d'),'<meta property="dcterms:modified">%s</meta>'%agora.strftime('%Y-%m-%dT%H:%M:%SZ'),
          '<meta property="schema:accessMode">textual</meta>']+(['<meta property="schema:accessMode">visual</meta>'] if tem_img else [])+[
          '<meta property="schema:accessModeSufficient">textual</meta>','<meta property="schema:accessibilityFeature">structuralNavigation</meta>',
          '<meta property="schema:accessibilityFeature">tableOfContents</meta>']+(['<meta property="schema:accessibilityFeature">alternativeText</meta>'] if alt_ok else [])+[
          '<meta property="schema:accessibilityFeature">readingOrder</meta>','<meta property="schema:accessibilityHazard">none</meta>',
          '<meta property="schema:conformsTo">https://www.w3.org/TR/WCAG22/</meta>','<meta property="schema:conformsTo">https://www.w3.org/TR/epub-a11y-11/</meta>',
          '<meta property="schema:accessibilitySummary">Conteúdo em português do Brasil, com estrutura de títulos, sumário navegável, tabelas com cabeçalhos%s e fonte adequada para dislexia (%s).</meta>'%(', imagens com texto alternativo' if alt_ok else '',E(nome_f))]
    opf=('<?xml version="1.0" encoding="utf-8"?>\n<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="pub-id" xml:lang="pt-BR" prefix="schema: http://schema.org/">\n'
         '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">\n%s\n</metadata>\n<manifest>\n%s\n</manifest>\n<spine toc="ncx">\n%s\n</spine>\n</package>\n')%('\n'.join(meta),'\n'.join(man),'\n'.join(spine))
    po=[1]
    def np_(rotulo,src,filhos=''):
        n=po[0];po[0]+=1
        return '<navPoint id="np%d" playOrder="%d"><navLabel><text>%s</text></navLabel><content src="%s"/>%s</navPoint>'%(n,n,E(rotulo),src,filhos)
    pts=[np_(titulo,'capa.xhtml')]
    for nome,sid,tit,h3 in toc:
        cab=np_(tit,nome+('#'+sid if sid else ''))
        filhos=''.join(np_(t_,'%s#%s'%(nome,i_)) for i_,t_ in h3)
        pts.append(cab[:-len('</navPoint>')]+filhos+'</navPoint>')
    prof=2 if any(h3 for _,_,_,h3 in toc) else 1
    ncx=('<?xml version="1.0" encoding="utf-8"?>\n<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1" xml:lang="pt-BR">\n<head>\n'
         '<meta name="dtb:uid" content="%s"/>\n<meta name="dtb:depth" content="%d"/>\n<meta name="dtb:totalPageCount" content="0"/>\n<meta name="dtb:maxPageNumber" content="0"/>\n</head>\n'
         '<docTitle><text>%s</text></docTitle>\n<navMap>\n%s\n</navMap>\n</ncx>\n')%(ident,prof,E(titulo_opf),'\n'.join(pts))
    container=('<?xml version="1.0" encoding="utf-8"?>\n<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">\n'
               '<rootfiles>\n<rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>\n</rootfiles>\n</container>\n')
    for x in(opf,ncx,container):etree.fromstring(x.encode('utf-8'))
    with zipfile.ZipFile(path,'w') as z:
        DT=(2026,1,1,0,0,0)
        def put(nm,dados,comp=zipfile.ZIP_DEFLATED):
            zi=zipfile.ZipInfo(nm,DT);zi.compress_type=comp;z.writestr(zi,dados)
        put('mimetype','application/epub+zip',zipfile.ZIP_STORED)  # primeiro item, sem compressão e sem nova linha
        put('META-INF/container.xml',container);put('OEBPS/content.opf',opf);put('OEBPS/toc.ncx',ncx);put('OEBPS/css/style.css',css)
        put('OEBPS/nav.xhtml',docs['nav.xhtml'])
        for _,nome,_,_,_ in arquivos:put('OEBPS/'+nome,docs[nome])
        for peso,estilo,arq,cam in faces:put('OEBPS/fonts/'+arq,open(cam,'rb').read())
        if faces and os.path.isfile(os.path.join(FONTES_DIR,ofl)):put('OEBPS/fonts/'+ofl,open(os.path.join(FONTES_DIR,ofl),'rb').read())
        for nm,mime,raw in cx.imgs.values():put('OEBPS/'+nm,raw)
    return path
# ============================== PPTX ==============================
def slides_pptx(t,sl,path,foot,tema=None,audios=None,avancar=False,duracoes=None):
    T=tema_de(tema);H=lambda h:thex(h,T);FONTE=T['pptx_font']
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    C=lambda h:RGBColor.from_string(thex(h,T))
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.enum.lang import MSO_LANGUAGE_ID
    PT_BR=MSO_LANGUAGE_ID.BRAZILIAN_PORTUGUESE
    MIN_CORPO=16  # leitura acessível (1.66): corpo nunca abaixo de 16 pt; tabelas, código e legendas no mínimo 12 pt
    def alt_pic(pic,alt,decorativa=False):
        "Texto alternativo da imagem; fundo do modelo é marcado como decorativo (descr vazio)."
        c=pic._element.nvPicPr.cNvPr;c.set('descr',alt or '');c.set('title',alt or '')
        if decorativa:
            from pptx.oxml import parse_xml
            c.append(parse_xml('<a:extLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:ext uri="{C183D7F6-B498-43B3-948B-1728B52AA6E4}">'
                               '<adec:decorative xmlns:adec="http://schemas.microsoft.com/office/drawing/2017/decorative" val="1"/></a:ext></a:extLst>'))
        elif not alt:aviso('Imagem sem descrição (texto alternativo vazio) na apresentação PPTX.')
    def esp(par,linha,depois):par.line_spacing=linha;par.space_after=Pt(depois)
    PRI=C('0B3D6E');ACC=C('0A7CC1');AC2=C('19A974');INK=C('1B2430')
    WH=C('FFFFFF');MU=C('5A6676');SOFT=C('EEF4FA');DEST=RGBColor.from_string(T['pptx_destaque']) if T.get('pptx_destaque') else AC2
    modelo=T.get('pptx_modelo'); escuro=bool(T.get('pptx_escuro'))
    if modelo:
        # Paleta extraída dos arquivos oficiais de referência ANAC 20 anos.
        PRI=RGBColor.from_string('FFFFFF' if escuro else '163B57');ACC=RGBColor.from_string('D5A900');AC2=RGBColor.from_string('16B8C4')
        INK=RGBColor.from_string('FFFFFF' if escuro else '1F3545');WH=RGBColor.from_string('FFFFFF');MU=RGBColor.from_string('DCE8EE' if escuro else '607786');SOFT=RGBColor.from_string('14334A' if escuro else 'EFF7FA')
    P=Presentation();P.slide_width=Inches(13.333);P.slide_height=Inches(7.5);W=13.333;L=0.62;CW=W-2*L
    pasta_modelo=os.path.join(os.path.dirname(os.path.abspath(__file__)),'temas',modelo) if modelo else ''
    def fundo_anac(s,idx,cls):
        if not modelo:return False
        if cls=='title' and idx==0:nome='capa.png'
        elif cls=='title' and idx==len(sl)-1:nome='encerramento.png'
        else:nome='conteudo-%s.png'%('a' if idx%2 else 'b')
        cam=os.path.join(pasta_modelo,nome)
        if os.path.isfile(cam):alt_pic(s.shapes.add_picture(cam,0,0,width=Inches(W),height=Inches(7.5)),'',True);return True
        return False
    def rect(s,x,y,w,h,col,shape=MSO_SHAPE.RECTANGLE):
        r=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h));r.fill.solid();r.fill.fore_color.rgb=col;r.line.fill.background();r.shadow.inherit=False;return r
    def addruns(par,t,size,color,bold=False):
        for x,b,i,c,u in runs(t):
            if c==2:x='[%s]'%(x or 'imagem');u=None;c=0
            r=par.add_run();r.text=x;f=r.font;f.size=Pt(size);f.bold=bool(b) or bold;f.italic=bool(i);f.name='Consolas' if c else FONTE;f.language_id=PT_BR
            f.color.rgb=ACC if (u and color==INK) else (PRI if (b and color==INK) else color)
            if u:r.hyperlink.address=u;f.underline=True
    def tbox(s,x,y,w,h):
        tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame;tf.word_wrap=True
        for a in('margin_left','margin_right','margin_top','margin_bottom'):setattr(tf,a,Inches(0.04))
        return tb,tf
    def lines(t,size,width):
        cpl=max(10,int(width*72/(size*0.60)));return max(1,-(-len(plain(t))//cpl))
    def est(b,size):
        k=b[0];lh=size*1.42/72
        if k=='p':return lines(b[1],size,CW)*lh+0.08
        if k=='list':return sum(lines(x[3],size-2*min(x[0],1),CW-0.4-0.3*x[0])*lh+0.08 for x in b[1])+0.12
        if k=='table':
            nc=len(b[1][0]);cw=CW/nc;ts=size-3
            return sum(max(lines(c,ts,cw-0.2) for c in r)*ts*1.4/72+0.16 for r in b[1])+0.18
        if k=='quote':return sum(est(x,size-1) for x in b[2])+0.32
        if k=='code':return est(('table',mermaid_edges(b[2])),size) if b[1]=='mermaid' else len(b[2].split('\n'))*0.22+0.2
        if k=='h':return 0.5
        if k=='img':
            raw=img_bytes(b[2])
            if raw is None:return 0.4
            w,h=img_size(raw);hh=min(4.6,CW*h/max(w,1),h/96)
            return hh+(0.35 if b[1] else 0)+0.1
        return 0.4
    def draw(s,b,y,size,col,x=L,w=CW):
        k=b[0];h=est(b,size)
        if k=='p':
            tb,tf=tbox(s,x,y,w,h);esp(tf.paragraphs[0],1.2,6);addruns(tf.paragraphs[0],b[1],size,col)
        elif k=='h':
            tb,tf=tbox(s,x,y,w,h);esp(tf.paragraphs[0],1.1,4);addruns(tf.paragraphs[0],b[2],size+4,ACC if col==INK else col,True)
        elif k=='list':
            tb,tf=tbox(s,x,y,w,h);n=0
            for dp,od,ck,t in b[1]:
                p=tf.paragraphs[0] if n==0 else tf.add_paragraph();n+=1
                esp(p,1.2,6);sz=max(MIN_CORPO,size-2*min(dp,1))
                bullet=('☑ ' if ck else '☐ ') if ck is not None else ('%d. '%n if od and dp==0 else '• ' if dp==0 else '– ')
                r=p.add_run();r.text='    '*dp+bullet;r.font.size=Pt(sz);r.font.bold=True;r.font.color.rgb=ACC if col==INK else col;r.font.name=FONTE;r.font.language_id=PT_BR
                addruns(p,t,sz,col)
        elif k=='table':
            rows=b[1];nc=max(len(r) for r in rows);ts=max(12,size-3)
            gt=s.shapes.add_table(len(rows),nc,Inches(x),Inches(y),Inches(w),Inches(h-0.1)).table
            gt.first_row=True;gt.horz_banding=True  # cabeçalho estrutural (firstRow), não só visual
            for ri,r in enumerate(rows):
                for ci in range(nc):
                    c=gt.cell(ri,ci);c.margin_left=c.margin_right=Inches(0.08);c.margin_top=c.margin_bottom=Inches(0.04)
                    tf=c.text_frame;tf.word_wrap=True;tf.paragraphs[0].text='';esp(tf.paragraphs[0],1.1,0)
                    addruns(tf.paragraphs[0],r[ci] if ci<len(r) else '',ts,WH if ri==0 else INK,ri==0)
                    c.fill.solid();c.fill.fore_color.rgb=PRI if ri==0 else (SOFT if ri%2==0 else WH)
        elif k=='quote':
            bg,bd={'info':(C('EAF1F8'),C('0A7CC1')),'tip':(C('E9F6F0'),C('1E7B4F')),'warning':(C('FFF4E0'),C('9A5B00'))}.get(b[1],(C('E9F6F0'),AC2))
            rect(s,x,y,w,h-0.05,bg);rect(s,x,y,0.08,h-0.05,bd)
            yy=y+0.1
            for q in b[2]:draw(s,q,yy,max(MIN_CORPO,size-1),col,x+0.2,w-0.3);yy+=est(q,max(MIN_CORPO,size-1))
        elif k=='img':
            raw=img_bytes(b[2])
            if raw is None:
                tb,tf=tbox(s,x,y,w,0.4);addruns(tf.paragraphs[0],'[imagem: %s]'%(b[1] or 'sem descrição'),size,col)
            else:
                iw,ih=img_size(raw);hh=min(4.6,w*ih/max(iw,1),ih/96);ww=hh*iw/max(ih,1)
                alt_pic(s.shapes.add_picture(io.BytesIO(raw),Inches(x+(w-ww)/2),Inches(y),Inches(ww),Inches(hh)),b[1])
                if b[1]:
                    tb,tf=tbox(s,x,y+hh+0.03,w,0.32);tf.paragraphs[0].alignment=2;addruns(tf.paragraphs[0],b[1],12,MU)
        elif k=='code':
            if b[1]=='mermaid':draw(s,('table',mermaid_edges(b[2])),y,size,col,x,w)
            else:
                tb,tf=tbox(s,x,y,w,h);r=tf.paragraphs[0].add_run();r.text=b[2];r.font.name='Consolas';r.font.size=Pt(12);r.font.color.rgb=col;r.font.language_id=PT_BR
        return h
    for idx,(cls,bl,notes) in enumerate(sl):
        s=P.slides.add_slide(P.slide_layouts[6]);istitle=cls=='title';tem_fundo=fundo_anac(s,idx,cls)
        if not tem_fundo:
            if istitle:rect(s,0,0,W,7.5,PRI);rect(s,0,0,0.15,7.5,DEST)
            else:
                if cls=='highlight':rect(s,0,0,W,7.5,C('F2F7FC'))
                rect(s,0,0,0.15,7.5,PRI)
        h1=[b for b in bl if b[0]=='h' and b[1]==1];rest=[b for b in bl if not(b[0]=='h' and b[1]==1)]
        if istitle:
            sub=[b for b in rest if b[0]=='h'];body=[b for b in rest if b[0]!='h'];y=2.2
            if modelo and idx==0:L0=5.15;CW0=7.35
            elif modelo and idx==len(sl)-1:L0=0.85;CW0=11.63;y=4.65
            else:L0=L;CW0=CW
            if h1:
                tb,tf=tbox(s,L0,y,CW0,1.3);tf.vertical_anchor=MSO_ANCHOR.BOTTOM;esp(tf.paragraphs[0],1.1,0);addruns(tf.paragraphs[0],h1[0][2],50,(WH if not modelo or escuro else PRI),True);y+=1.4
                rect(s,L0,y,min(CW0,len(plain(h1[0][2]))*0.36),0.06,DEST);y+=0.25
            for b in sub:tb,tf=tbox(s,L0,y,CW0,0.7);addruns(tf.paragraphs[0],b[2],28,(C('D6ECFF') if not modelo else (WH if escuro else INK)),True);y+=0.75
            for b in body:y+=draw(s,b,y,20,(WH if escuro else INK) if modelo else C('E8F2FB'),L0,CW0)
        else:
            y=0.62 if modelo else 0.45
            if h1:
                tb,tf=tbox(s,L,y,CW,0.9);tf.vertical_anchor=MSO_ANCHOR.BOTTOM;esp(tf.paragraphs[0],1.1,0);addruns(tf.paragraphs[0],h1[0][2],34,PRI,True);y+=0.95
                rect(s,L,y,CW,0.05,ACC);y+=0.25
            avail=6.75-y;size=22
            while size>MIN_CORPO and sum(est(b,size) for b in rest)>avail:size-=1  # se não couber em 16 pt, o texto flui (não volta a 12 pt)
            for b in rest:y+=draw(s,b,y,size,INK)
        tb,tf=tbox(s,L,7.05,8,0.3);addruns(tf.paragraphs[0],foot,11,(WH if modelo and escuro else MU) if modelo else (C('CFE3F5') if istitle else MU))
        tb,tf=tbox(s,W-1.5,7.05,1,0.3);addruns(tf.paragraphs[0],'%d / %d'%(idx+1,len(sl)),11,C('CFE3F5') if istitle else MU)
        if notes:
            ntf=s.notes_slide.notes_text_frame;ntf.text=notes
            for pp in ntf.paragraphs:
                for rr in pp.runs:rr.font.language_id=PT_BR
        if audios and idx in audios:
            try:_pptx_audio(s,audios[idx],W,avancar,(duracoes or {}).get(idx))
            except Exception as e:print('AVISO: narração do slide %d não incluída no PPTX: %s'%(idx+1,e))
    P.core_properties.title=t;P.core_properties.language='pt-BR';P.save(path)
def _icone_som():
    try:
        from PIL import Image,ImageDraw
        im=Image.new('RGBA',(96,96),(0,0,0,0));d=ImageDraw.Draw(im)
        d.rounded_rectangle((2,2,94,94),18,fill=(19,81,180,255));d.polygon([(22,38),(38,38),(56,22),(56,74),(38,58),(22,58)],fill='white')
        d.arc((50,30,78,66),-50,50,fill='white',width=5);d.arc((52,20,90,76),-50,50,fill='white',width=5)
        b=io.BytesIO();im.save(b,'PNG');return b.getvalue()
    except Exception:
        return base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==')
def _wav_dur(dados):
    try:
        import struct
        if dados[:4]==b'RIFF':
            p=12;br=0
            while p+8<=len(dados):
                cid=dados[p:p+4];tam=struct.unpack('<I',dados[p+4:p+8])[0]
                if cid==b'fmt ':br=struct.unpack('<I',dados[p+16:p+20])[0]
                if cid==b'data' and br:return tam/br
                p+=8+tam+(tam&1)
    except Exception:pass
    return None
def _mp3_dur(dados):
    """Duração aproximada de um MP3 (CBR) pelo cabeçalho do primeiro quadro."""
    try:
        i=0
        if dados[:3]==b'ID3':i=10+((dados[6]&127)<<21|(dados[7]&127)<<14|(dados[8]&127)<<7|(dados[9]&127))
        while i<len(dados)-4 and not(dados[i]==0xFF and dados[i+1]&0xE0==0xE0):i+=1
        h=dados[i+1:i+3];ver=(h[0]>>3)&3;br_i=(h[1]>>4)&15
        tab1=[0,32,40,48,56,64,80,96,112,128,160,192,224,256,320];tab2=[0,8,16,24,32,40,48,56,64,80,96,112,128,144,160]
        br=(tab1 if ver==3 else tab2)[br_i]*1000
        return (len(dados)-i)*8/br if br else None
    except Exception:return None
def _pptx_audio(s,cam,W,avancar,dur=None):
    """Embute o áudio no slide (início automático, ícone fora da área visível) e, se pedido, avança o slide ao fim."""
    from pptx.media import Video
    from pptx.oxml import parse_xml
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.util import Inches
    ext=os.path.splitext(cam)[1].lower();mime=MIME_AUDIO.get(ext,'audio/mpeg');dados=open(cam,'rb').read()
    part=s.part;media=Video.from_blob(dados,mime,'narracao'+ext)
    mp=part.package.get_or_add_media_part(media)
    r_media=part.relate_to(mp,RT.MEDIA);r_audio=part.relate_to(mp,RT.AUDIO)
    _,r_img=part.get_or_add_image_part(io.BytesIO(_icone_som()))
    tree=s.shapes._spTree;sid=max([int(x) for x in tree.xpath('//p:cNvPr/@id')]+[1])+1
    x,y,tam=Inches(W+0.3),Inches(0.3),Inches(0.4)
    pic=parse_xml(('<p:pic xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
      '<p:nvPicPr><p:cNvPr id="%d" name="Narração" descr="Narração do slide"><a:hlinkClick r:id="" action="ppaction://media"/></p:cNvPr><p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr>'
      '<p:nvPr><a:audioFile r:link="%s"/><p:extLst><p:ext uri="{DAA4B4D4-6D71-4841-9C94-3DE7FCFB9230}"><p14:media xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" r:embed="%s"/></p:ext></p:extLst></p:nvPr></p:nvPicPr>'
      '<p:blipFill><a:blip r:embed="%s"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
      '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>')%(sid,r_audio,r_media,r_img,x,y,tam,tam))
    tree.append(pic)
    sld=s._element
    for velho in sld.xpath('./p:transition|./p:timing'):sld.remove(velho)
    ns='xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    novos=[]
    dur=dur or (_wav_dur(dados) if ext=='.wav' else _mp3_dur(dados) if ext=='.mp3' else None)
    if avancar and dur:novos.append(parse_xml('<p:transition %s advTm="%d"/>'%(ns,int(dur*1000)+800)))
    novos.append(parse_xml(('<p:timing %s><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
      '<p:par><p:cTn id="2" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
      '<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
      '<p:par><p:cTn id="4" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
      '<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="5" dur="%d" fill="hold"/><p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cBhvr></p:cmd>'
      '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
      '<p:audio><p:cMediaNode vol="80000"><p:cTn id="6" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
      '<p:endCondLst><p:cond evt="onStopAudio" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:endCondLst></p:cTn><p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cMediaNode></p:audio>'
      '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')%(ns,int((dur or 1)*1000),sid,sid)))
    ext_=sld.xpath('./p:extLst')
    for el in novos:
        if ext_:ext_[0].addprevious(el)
        else:sld.append(el)
# ============================== EXECUÇÃO ==============================
def markdown_com_tema(md,tema):
    """Insere ou atualiza `theme` no frontmatter YAML sem alterar o restante do Markdown."""
    md=md or '';tema=tema if tema in TEMAS else TEMA_PADRAO
    if md.startswith('---\n'):
        fim=md.find('\n---',4)
        if fim!=-1:
            cab=md[4:fim]
            if re.search(r'(?m)^theme\s*:',cab):cab=re.sub(r'(?m)^theme\s*:.*$','theme: '+tema,cab)
            else:cab=cab.rstrip()+'\ntheme: '+tema
            return '---\n'+cab+'\n---'+md[fim+4:]
    return '---\ntheme: '+tema+'\n---\n\n'+md.lstrip()
def gerar(gm,sm,prefixo='Conteudo',rodape='',barra='',saida='.',tema=None,narracao=None,fluxos=None):
    tema=tema if tema in TEMAS else TEMA_PADRAO
    """Gera os arquivos do que estiver preenchido: apostila (MD/HTML/DOCX/EPUB/EPUB OpenDyslexic) e/ou apresentação (MD/HTML/PPTX)."""
    global _FONTES_CSS
    del AVISOS[:];_FONTES_CSS=None  # relê as fontes a cada geração (o app mantém o módulo carregado)
    gm=gm or '';sm=sm or ''
    tem_g,tem_s=bool(gm.strip()),bool(sm.strip())
    if not tem_g and not tem_s:raise ValueError('Nada para gerar: a apostila e a apresentação estão vazias.')
    os.makedirs(saida,exist_ok=True)
    P=lambda e,n:os.path.join(saida,n+prefixo+e)
    arqs=[];msg='';g='';sl=[]
    if tem_g:
        gm_=sem_comentarios(gm);g=treinamento_html(gm_,barra,tema)
        open(P('.md','Apostila_'),'w',encoding='utf-8').write(gm);open(P('.html','Apostila_'),'w',encoding='utf-8').write(g)
        arqs+=[P('.md','Apostila_'),P('.html','Apostila_')]
        try:treinamento_docx(gm_,P('.docx','Apostila_'),rodape or title_of(gm_,''),tema);arqs.append(P('.docx','Apostila_'));msg+=' | docx: ok'
        except Exception as e:msg+=' | docx: ERRO %s'%e
        ne=0
        for suf,fonte in(('.epub','atkinson'),('_OpenDyslexic.epub','opendyslexic')):
            try:treinamento_epub(gm_,P(suf,'Apostila_'),rodape or title_of(gm_,''),tema,fonte);arqs.append(P(suf,'Apostila_'));ne+=1
            except Exception as e:msg+=' | epub %s: ERRO %s'%(fonte,e)
        msg+=' | epub: %d'%ne
    else:msg+=' | apostila: vazia (não gerada)'
    if tem_s:
        sm=markdown_com_tema(sm,tema);t,sl=split_slides(sm);nr=_narr(narracao)
        open(P('.md','Apresentacao_'),'w',encoding='utf-8').write(sm);open(P('.html','Apresentacao_'),'w',encoding='utf-8').write(slides_html(t,sl,rodape,tema,nr['html'],nr['avancar'],_fluxos_dict(fluxos)))
        arqs+=[P('.md','Apresentacao_'),P('.html','Apresentacao_')]
        try:slides_pptx(t,sl,P('.pptx','Apresentacao_'),rodape,tema,nr['pptx'],nr['avancar'],nr['duracoes']);arqs.append(P('.pptx','Apresentacao_'));msg+=' | pptx: ok'
        except Exception as e:msg+=' | pptx: ERRO %s'%e
        if nr['html'] or nr['pptx']:msg+=' | narração: %d slide(s)'%len(set(nr['html'])|set(nr['pptx']))
    else:msg+=' | apresentação: vazia (não gerada)'
    res='OK | slides: %d | checklists: %d | exercicios: %d%s | links: %d | imagens: %d'%(len(sl),g.count('class="task"'),len(re.findall(r'<h3 id="exercicio',g)),msg,g.count('<a href="http'),len(IMG.findall(CODIGO_RX.sub('',gm)))+len(IMG.findall(CODIGO_RX.sub('',sm))))+' | tema: '+tema
    if AVISOS:res+=' | avisos: '+' / '.join(AVISOS[:6])+(' (+%d)'%(len(AVISOS)-6) if len(AVISOS)>6 else '')
    return arqs,res
if __name__=='__main__':
    a=argparse.ArgumentParser();a.add_argument('apostila',help="arquivo da apostila ('-' para não gerar)");a.add_argument('apresentacao',nargs='?',default='-',help="arquivo da apresentação ('-' ou omitido para não gerar)");a.add_argument('--prefixo',default='Curso');a.add_argument('--rodape',default='');a.add_argument('--barra',default='');a.add_argument('--saida',default='.')
    a.add_argument('--sem-embutir',action='store_true',help='não converter imagens para base64');a.add_argument('--tema',default=TEMA_PADRAO,choices=sorted(TEMAS),help='layout dos arquivos (padrão: %(default)s)');a.add_argument('--audios',default='',help='pasta com a narração por slide: 1.mp3, 2.ogg, 3.wav… (número do slide)');a.add_argument('--avancar',action='store_true',help='avançar os slides ao fim da narração');a=a.parse_args()
    ler=lambda f:'' if f in('-','') else open(f,encoding='utf-8').read()
    gm=ler(a.apostila);sm=ler(a.apresentacao)
    if not a.sem_embutir:
        gm,n1,f1=embed_images(gm,os.path.dirname(os.path.abspath(a.apostila if a.apostila!='-' else '.')));sm,n2,f2=embed_images(sm,os.path.dirname(os.path.abspath(a.apresentacao if a.apresentacao!='-' else '.')))
        if n1+n2:print('Imagens embutidas em base64: %d'%(n1+n2))
        for f in f1+f2:print('AVISO: imagem não encontrada/baixada: %s'%f)
    narr=None
    if a.audios and os.path.isdir(a.audios):
        narr={'html':{},'pptx':{},'avancar':a.avancar}
        for f in sorted(os.listdir(a.audios)):
            m=re.match(r'^(\d+)\.(mp3|ogg|wav)$',f,re.I)
            if not m:continue
            k=str(int(m.group(1))-1);cam=os.path.join(a.audios,f)
            narr['html'][k]=cam
            if m.group(2).lower()!='ogg' and (k not in narr['pptx'] or m.group(2).lower()=='mp3'):narr['pptx'][k]=cam
    try:_,res=gerar(gm,sm,a.prefixo,a.rodape,a.barra,a.saida,a.tema,narr)
    except ValueError as e:print('ERRO: %s'%e);raise SystemExit(1)
    print(res)
    if 'ERRO' in res:raise SystemExit(1)
